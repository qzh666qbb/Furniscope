"""Secure product-document text extraction and structured attribute contract."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path
from typing import Any
import re
import zipfile

import fitz
from openpyxl import load_workbook
from PIL import Image
from pydantic import BaseModel, ConfigDict, Field
import pytesseract


class DocumentParseError(ValueError):
    def __init__(self, code: str, message: str) -> None:
        super().__init__(message)
        self.code = code


class ExtractedAttribute(BaseModel):
    attribute_code: str = Field(pattern=r"^[a-z][a-z0-9_]{1,99}$")
    value: Any
    unit: str | None = Field(default=None, max_length=32)
    confidence: float = Field(ge=0, le=1)
    evidence_text: str = Field(min_length=1, max_length=1000)
    model_config = ConfigDict(extra="forbid")


class ProductDocumentExtraction(BaseModel):
    product_name: str | None = Field(default=None, max_length=200)
    sku: str | None = Field(default=None, max_length=100)
    attributes: list[ExtractedAttribute] = Field(min_length=1, max_length=50)
    model_config = ConfigDict(extra="forbid")


class CatalogItem(BaseModel):
    sku: str
    category: str | None = None
    page: int = Field(ge=1)


class ParsedDocument(BaseModel):
    filename: str
    mime_type: str
    text: str
    extraction_method: str
    page_or_sheet_count: int = Field(ge=1)
    catalog_items: list[CatalogItem] = Field(default_factory=list)


class DocumentParser:
    allowed = {
        "application/pdf": {".pdf"},
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": {".xlsx"},
        "image/jpeg": {".jpg", ".jpeg"},
        "image/png": {".png"},
    }

    def __init__(self, *, max_bytes: int, max_pages: int = 80, max_text_chars: int = 120_000) -> None:
        self.max_bytes = max_bytes
        self.max_pages = max_pages
        self.max_text_chars = max_text_chars

    def parse(self, filename: str, mime_type: str, content: bytes) -> ParsedDocument:
        suffix = Path(filename).suffix.lower()
        if mime_type not in self.allowed or suffix not in self.allowed[mime_type]:
            raise DocumentParseError("FILE_TYPE_UNSUPPORTED", "仅支持 PDF、XLSX、JPG/JPEG 和 PNG")
        if not content or len(content) > self.max_bytes:
            raise DocumentParseError("FILE_SIZE_INVALID", "文件为空或超过上传大小限制")
        if mime_type == "application/pdf":
            return self._pdf(filename, mime_type, content)
        if mime_type.endswith("spreadsheetml.sheet"):
            return self._xlsx(filename, mime_type, content)
        return self._image(filename, mime_type, content)

    def _pdf(self, filename: str, mime_type: str, content: bytes) -> ParsedDocument:
        if not content.startswith(b"%PDF-"):
            raise DocumentParseError("FILE_SIGNATURE_INVALID", "PDF 文件签名无效")
        try:
            document = fitz.open(stream=content, filetype="pdf")
        except Exception as exc:
            raise DocumentParseError("PDF_PARSE_FAILED", "PDF 无法打开或已损坏") from exc
        if document.page_count < 1 or document.page_count > self.max_pages:
            document.close()
            raise DocumentParseError("PDF_PAGE_LIMIT", "PDF 页数为空或超过限制")
        page_count = document.page_count
        parts: list[str] = []
        catalog_items: list[CatalogItem] = []
        seen_skus: set[str] = set()
        used_ocr = False
        try:
            for page in document:
                text = page.get_text("text").strip()
                if not text:
                    used_ocr = True
                    pixmap = page.get_pixmap(matrix=fitz.Matrix(1.8, 1.8), alpha=False)
                    image = Image.open(BytesIO(pixmap.tobytes("png")))
                    text = pytesseract.image_to_string(image, lang="chi_sim+eng").strip()
                if text:
                    page_number = page.number + 1
                    parts.append(f"[page:{page_number}]\n{text}")
                    lines = [line.strip() for line in text.splitlines() if line.strip()]
                    category = next((line for line in lines
                                     if line == line.upper() and 4 <= len(line) <= 80
                                     and "FURNITURE" not in line and not line.startswith("HF-")), None)
                    for sku in dict.fromkeys(re.findall(r"\bHF-[A-Z][0-9A-Z-]{2,}\b", text.upper())):
                        if sku in seen_skus:
                            continue
                        seen_skus.add(sku)
                        catalog_items.append(CatalogItem(sku=sku, category=category, page=page_number))
        finally:
            document.close()
        text = "\n".join(parts)[:self.max_text_chars]
        if not text:
            raise DocumentParseError("DOCUMENT_TEXT_EMPTY", "PDF 文本与 OCR 均未提取到内容")
        return ParsedDocument(filename=filename, mime_type=mime_type, text=text,
                              extraction_method="pdf_text+ocr" if used_ocr else "pdf_text",
                              page_or_sheet_count=page_count,
                              catalog_items=catalog_items)

    def _xlsx(self, filename: str, mime_type: str, content: bytes) -> ParsedDocument:
        if not zipfile.is_zipfile(BytesIO(content)):
            raise DocumentParseError("FILE_SIGNATURE_INVALID", "XLSX 文件签名无效")
        try:
            workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
            lines: list[str] = []
            for sheet in workbook.worksheets:
                lines.append(f"[sheet:{sheet.title}]")
                for row in sheet.iter_rows(values_only=True):
                    values = [str(value).strip() for value in row if value not in (None, "")]
                    if values:
                        lines.append(" | ".join(values))
                    if sum(map(len, lines)) >= self.max_text_chars:
                        break
            count = len(workbook.worksheets)
            workbook.close()
        except Exception as exc:
            raise DocumentParseError("SPREADSHEET_PARSE_FAILED", "XLSX 无法读取或已损坏") from exc
        text = "\n".join(lines)[:self.max_text_chars]
        if not text:
            raise DocumentParseError("DOCUMENT_TEXT_EMPTY", "XLSX 中没有可提取内容")
        return ParsedDocument(filename=filename, mime_type=mime_type, text=text,
                              extraction_method="xlsx_cells", page_or_sheet_count=count)

    def _image(self, filename: str, mime_type: str, content: bytes) -> ParsedDocument:
        try:
            image = Image.open(BytesIO(content))
            image.verify()
            image = Image.open(BytesIO(content))
            text = pytesseract.image_to_string(image, lang="chi_sim+eng").strip()
        except Exception as exc:
            raise DocumentParseError("IMAGE_PARSE_FAILED", "图片无法读取或 OCR 失败") from exc
        if not text:
            raise DocumentParseError("IMAGE_VISION_REQUIRED", "图片无可识别文字，需要配置视觉理解模型")
        return ParsedDocument(filename=filename, mime_type=mime_type, text=text[:self.max_text_chars],
                              extraction_method="image_ocr", page_or_sheet_count=1)
