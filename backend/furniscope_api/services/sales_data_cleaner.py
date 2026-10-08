"""Deterministic intake: suggestions never change facts or bypass confirmation."""

from __future__ import annotations

import hashlib
import csv
import io
import json
import math
import re
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any

import pandas as pd

from ..errors import BusinessError
from ..schemas.data_imports import ImportRules

RULE_VERSION = "sales-contract-v2"
MAX_ROWS = 100_000
ALIASES = {
    "date": {"date", "日期", "时间", "订单日期", "销售日期"},
    "sku": {"sku", "商品编码", "产品编码", "货号", "seller sku"},
    "site": {"site", "site_code", "站点", "市场", "国家"},
    "sales": {"sales", "daily_sales", "销量", "销售数量", "quantity", "数量"},
    "inventory": {"inventory", "库存", "库存数量"},
    "status": {"status", "销售状态", "生命周期"},
    "warehouse": {"warehouse", "仓库", "仓库代码"},
    "order_id": {"order_id", "订单号"},
    "line_id": {"line_id", "订单行号"},
    "order_status": {"order_status", "订单状态"},
    "refunded_units": {"refunded_units", "退货件数", "退款件数"},
    "unit_price": {"unit_price", "成交单价"},
    "discount": {"discount", "折扣比例"},
    "currency": {"currency", "币种"},
}


def stable_json(value: Any) -> bytes:
    return json.dumps(value, ensure_ascii=False, sort_keys=True,
                      separators=(",", ":"), allow_nan=False).encode("utf-8")


def sha256(content: bytes) -> str:
    return hashlib.sha256(content).hexdigest()


def read_table(filename: str, content: bytes) -> pd.DataFrame:
    if not content:
        raise BusinessError("FILE_EMPTY", "上传文件为空", status_code=400)
    try:
        suffix = Path(filename).suffix.lower()
        header = None
        if suffix == ".csv":
            header = next(csv.reader(io.StringIO(content.decode("utf-8-sig"))))
            frame = pd.read_csv(io.BytesIO(content), dtype=str, keep_default_na=False,
                                nrows=MAX_ROWS + 1, encoding="utf-8-sig")
        elif suffix == ".xlsx":
            from openpyxl import load_workbook
            book = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            try:
                header = list(next(book.active.iter_rows(values_only=True)))
            finally:
                book.close()
            frame = pd.read_excel(io.BytesIO(content), dtype=object, keep_default_na=False,
                                  nrows=MAX_ROWS + 1, engine="openpyxl")
        elif suffix == ".json":
            def unique_object(pairs):
                result = {}
                for key, value in pairs:
                    if key in result:
                        raise ValueError("JSON字段重复，请先明确列名")
                    result[key] = value
                return result
            records = json.loads(content, object_pairs_hook=unique_object)
            if not isinstance(records, list) or not all(isinstance(r, dict) for r in records):
                raise ValueError("JSON 顶层必须是记录数组")
            frame = pd.DataFrame(records).fillna("")
        else:
            raise BusinessError("FILE_TYPE_UNSUPPORTED", "请上传 CSV、XLSX 或 JSON", status_code=415)
        if frame.empty or len(frame) > MAX_ROWS:
            raise ValueError(f"文件须包含 1—{MAX_ROWS} 行")
        columns = [str(c).strip() if c is not None else "" for c in
                   (header if header is not None else frame.columns)]
        if len(columns) != len(set(columns)) or any(not c for c in columns):
            raise ValueError("表头重复或为空，请先明确列名")
        frame.columns = columns
        return frame
    except BusinessError:
        raise
    except Exception as exc:
        raise BusinessError("DATA_IMPORT_INVALID", f"无法解析数据文件：{str(exc)[:200]}",
                            status_code=422) from exc


def suggest_mapping(columns: list[str]) -> dict[str, str]:
    # Monetary Amazon 'product sales' is deliberately NOT a unit-sales synonym.
    return {target: matches[0] for target, names in ALIASES.items()
            if len(matches := [c for c in columns if c.lower().strip() in names]) == 1}


def clean_table(frame: pd.DataFrame, rules: ImportRules) -> dict[str, Any]:
    mapping = rules.mapping
    value_field = rules.kind
    required = {"date", "sku", value_field} | (
        {"site"} if not rules.default_site and "warehouse" not in mapping else set())
    errors: list[dict] = []
    warnings: list[dict] = []
    if required - set(mapping) or set(mapping.values()) - set(frame.columns):
        raise BusinessError("DATA_MAPPING_INCOMPLETE", "请确认日期、SKU、站点和数量对应的列",
                            status_code=422)
    grouped: dict[tuple[str, str, str], dict] = {}
    seen_orders: dict[tuple, dict] = {}
    audit = []
    adjustments = {"duplicate_units": 0.0, "cancelled_units": 0.0, "refunded_units": 0.0}
    input_total = 0.0
    parsed_rows = 0
    for number, (_, raw) in enumerate(frame.iterrows(), 2):
        def field(name: str) -> Any:
            return raw[mapping[name]] if name in mapping else ""
        try:
            raw_date = field("date")
            if isinstance(raw_date, (date, datetime)):
                when = raw_date.date() if isinstance(raw_date, datetime) else raw_date
            else:
                when = datetime.strptime(str(raw_date).strip(), rules.date_format).date()
                # strptime accepts unpadded dates; ambiguity is removed by the explicit format.
            if not date(2000, 1, 1) <= when <= date.today():
                raise ValueError("日期不在2000年至今天的历史范围")
            sku = str(field("sku")).strip()
            site = str(field("site") or rules.default_site or "").strip().upper()
            warehouse = str(field("warehouse")).strip()
            if "warehouse" in mapping:
                warehouse_site = rules.warehouse_sites.get(warehouse, "").strip().upper()
                if not warehouse_site:
                    raise ValueError("仓库未映射到站点，不自动拆分共享库存")
                if site and site != warehouse_site:
                    raise ValueError("仓库映射与站点列/默认站点冲突")
                site = warehouse_site
            if not sku or len(sku) > 128:
                raise ValueError("SKU为空或超过128字符")
            if not re.fullmatch(r"[A-Z][A-Z0-9_-]{1,15}", site):
                raise ValueError("站点须为已确认的站点代码，例如 US/UK/DE")
            quantity = float(field(value_field))
            if not math.isfinite(quantity):
                raise ValueError("数量不是有限数值")
            if not math.isfinite(input_total + quantity):
                raise ValueError("数量汇总超出可表示范围，请核对数量口径")
            if quantity < 0 and (value_field == "inventory" or rules.sales_basis != "net_units"):
                raise ValueError("负数量须确认净销量口径；库存不能为负")
            status = str(field("status") or "active").strip().lower()
            if status not in {"active", "out_of_stock", "discontinued", "unknown"}:
                raise ValueError("销售状态须为 active/out_of_stock/discontinued/unknown")
            order_key = None
            refund = 0.0
            order_status = "completed"
            if "order_id" in mapping:
                order_id, line_id = str(field("order_id")).strip(), str(field("line_id")).strip()
                if not order_id or not line_id or max(len(order_id), len(line_id)) > 128:
                    raise ValueError("订单号/行号为空或超过128字符")
                order_key = (site, order_id, line_id)
                if quantity < 0:
                    raise ValueError("订单快照数量须非负，退货填写独立件数，避免二次扣减")
                if "order_status" in mapping:
                    order_status = rules.order_statuses.get(str(field("order_status")).strip())
                    if not order_status:
                        raise ValueError("订单状态未明确映射为 completed/cancelled/refunded")
                if "refunded_units" in mapping:
                    refund = float(field("refunded_units"))
                    if not math.isfinite(refund) or not 0 <= refund <= quantity:
                        raise ValueError("退货件数须非负且不超过原订单件数")
                elif order_status == "refunded":
                    raise ValueError("退款订单须填写实际退货件数；退款金额不代表整单退货")
            facts = {}
            for name in ("unit_price", "discount"):
                if name in mapping and str(field(name)).strip():
                    value = float(field(name))
                    if not math.isfinite(value) or value < 0 or name == "discount" and value > 1:
                        raise ValueError("成交单价须非负有限数，折扣须为0—1比例")
                    facts[name] = value
            currency = str(field("currency") or rules.default_currency or "").strip().upper()
            if "unit_price" in facts and not re.fullmatch(r"[A-Z]{3}", currency):
                raise ValueError("成交单价须有明确三位币种代码")
            if currency:
                if not re.fullmatch(r"[A-Z]{3}", currency):
                    raise ValueError("币种须为三位代码")
                facts["currency"] = currency
            normalized = dict(date=when.isoformat(), sku=sku, site=site, status=status,
                              quantity=quantity, refund=refund, order_status=order_status,
                              warehouse=warehouse, facts=facts)
            duplicate = seen_orders.get(order_key) if order_key else None
            if duplicate and (duplicate["value"] != normalized or rules.duplicate_orders == "error"):
                raise ValueError(f"订单号/行号重复或内容冲突（首次见于行{duplicate['row']}），请核对快照")
            effective = (0.0 if order_status == "cancelled" else
                         quantity - refund if rules.sales_basis == "net_units" else quantity)
            prior = grouped.get((when.isoformat(), sku, site))
            if prior and not math.isfinite(prior[value_field] + effective):
                raise ValueError("同日数量汇总超出可表示范围")
            if prior and facts.get("currency") and any(
                item.get("currency") not in (None, facts["currency"])
                for item in prior.get("price_facts", [])
            ):
                raise ValueError("同日SKU/站点存在混合币种，请先拆分或核对，不能汇率猜算")
        except (ValueError, TypeError, OverflowError) as exc:
            errors.append({"row": number, "reason": str(exc)})
            continue
        parsed_rows += 1
        input_total += quantity
        if duplicate:
            adjustments["duplicate_units"] += quantity
            audit.append({"row": number, "action": "duplicate_ignored", "original_row": duplicate["row"],
                          "input_units": quantity, "output_units": 0, "order_key": order_key})
            continue
        if order_key:
            seen_orders[order_key] = {"row": number, "value": normalized}
        if order_status == "cancelled":
            adjustments["cancelled_units"] += quantity
        elif rules.sales_basis == "net_units":
            adjustments["refunded_units"] += refund
        audit.append({"row": number, "action": "cancelled_excluded" if order_status == "cancelled" else "included",
                      "input_units": quantity, "output_units": effective, "refunded_units": refund,
                      "order_key": order_key, "warehouse": warehouse or None, "site": site,
                      "facts": facts})
        price_fact = {**facts, "source_row": number, "units": effective} if facts else None
        key = (when.isoformat(), sku, site)
        if key in grouped:
            if rules.grain == "daily":
                errors.append({"row": number, "reason": "日期/SKU/站点重复；日汇总不可重复"})
                continue
            if grouped[key]["status"] != status:
                errors.append({"row": number, "reason": "同日明细销售状态冲突"})
                continue
            grouped[key][value_field] += effective
            grouped[key]["source_rows"].append(number)
        else:
            grouped[key] = dict(date=key[0], sku=sku, site=site, status=status,
                                source_rows=[number], **{value_field: effective})
        if order_key:
            grouped[key].setdefault("order_keys", []).append(list(order_key))
        if price_fact:
            grouped[key].setdefault("price_facts", []).append(price_fact)
        # A differing transaction price is preserved per line; no invented daily average.
        for name in ("unit_price", "discount", "currency"):
            values = [item.get(name) for item in grouped[key].get("price_facts", [])]
            grouped[key][name] = values[0] if (
                values and len(values) == len(grouped[key]["source_rows"]) and
                all(value == values[0] for value in values)) else None

    records = sorted(grouped.values(), key=lambda r: (r["sku"], r["site"], r["date"]))
    pairs: dict[tuple[str, str], list[dict]] = {}
    for record in records:
        pairs.setdefault((record["sku"], record["site"]), []).append(record)
    missing_count = 0
    filled = 0
    for (sku, site), rows in pairs.items():
        first, last = date.fromisoformat(rows[0]["date"]), date.fromisoformat(rows[-1]["date"])
        existing = {r["date"] for r in rows}
        missing = [(first + timedelta(days=i)).isoformat() for i in range((last - first).days + 1)
                   if (first + timedelta(days=i)).isoformat() not in existing]
        missing_count += len(missing)
        if missing and rules.missing_dates == "zero" and value_field == "sales":
            if len(records) + len(missing) > MAX_ROWS:
                raise BusinessError("DATA_IMPORT_TOO_LARGE", "补零后超过10万行，请分批接入",
                                    status_code=422)
            records.extend(dict(date=day, sku=sku, site=site, sales=0.0, status="active",
                                source_rows=[], transformation="confirmed_missing_zero")
                           for day in missing)
            filled += len(missing)
        if missing:
            warnings.append({"code": "MISSING_DATES", "sku": sku, "site": site,
                             "count": len(missing), "examples": missing[:5]})
        values = [r[value_field] for r in rows]
        median = float(pd.Series(values).median())
        outliers = [r["source_rows"] for r in rows if r[value_field] > max(10 * median, 100)]
        if outliers:
            warnings.append({"code": "LARGE_VALUES_RETAINED", "sku": sku,
                             "count": len(outliers), "source_rows": outliers[:5]})
    negative = sum(r[value_field] < 0 for r in records)
    censored = sum(r["status"] != "active" for r in records)
    unresolved = []
    if missing_count > filled:
        unresolved.append("日期缺口为未知；补齐数据或确认完整导出后补零")
    if negative:
        unresolved.append("存在负净销量，须对账后提供非负需求标签；原值已保留")
    if censored:
        unresolved.append("存在缺货/停售/未知状态；需补充可观测需求数据")
    canonical = sorted(records, key=lambda r: (r["sku"], r["site"], r["date"]))
    total = sum(r[value_field] for r in canonical)
    if not math.isfinite(total):
        errors.append({"row": None, "reason": "标准数据汇总超出可表示范围"})
        total = None
    expected_total = input_total - sum(adjustments.values())
    reconciled = total is not None and math.isclose(expected_total, total, rel_tol=1e-9, abs_tol=1e-9)
    if not reconciled and not errors:
        errors.append({"row": None, "reason": "扣除重复/取消/退货后数量不平，请检查源文件"})
    report = {
        "rule_version": RULE_VERSION, "source_rows": len(frame), "parsed_rows": parsed_rows,
        "canonical_rows": len(canonical), "error_count": len(errors), "errors": errors[:100],
        "warnings": warnings[:100], "missing_dates": missing_count, "filled_rows": filled,
        "negative_rows": negative, "censored_rows": censored,
        "input_quantity_total": input_total, "output_quantity_total": total,
        "adjustments": adjustments, "expected_quantity_total": expected_total,
        "quantity_reconciled": reconciled,
        "price_fact_rows": sum(bool(item.get("facts")) for item in audit),
        "pair_count": len(pairs), "sku_count": len({r["sku"] for r in canonical}),
        "date_from": min((r["date"] for r in canonical), default=None),
        "date_to": max((r["date"] for r in canonical), default=None),
        "can_confirm": bool(canonical) and not errors,
        "trainable": bool(canonical) and not errors and not unresolved and value_field == "sales",
        "training_blockers": unresolved,
    }
    return {"records": canonical, "quality": report, "errors": errors, "row_audit": audit}
