"""Versioned product templates and all-or-nothing catalog imports."""

from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
import hashlib
import json
from pathlib import Path
import re
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncSession

from ..errors import BusinessError
from ..repositories.product_repository import ProductRepository
from .audit_service import AuditService
from .forecast_catalog import contract_context


TEMPLATE_VERSION = "product-master-2026.10.1"
MAX_FILE_SIZE = 20 * 1024 * 1024
MAX_SHEETS = 20
MAX_ROWS = 20_000
XLSX_MIME_TYPES = {
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "application/octet-stream",
    "application/zip",
}

FIELD_SCHEMA = [
    {"code": "sku", "title": "SKU", "required": True, "type": "text", "max_length": 100,
     "aliases": ["sku", "sku编码", "产品sku", "产品编码", "商品编码", "货号", "款号", "itemcode"]},
    {"code": "name", "title": "产品名称", "required": True, "type": "text", "max_length": 200,
     "aliases": ["name", "产品名称", "名称", "商品名称", "品名", "标题", "productname"]},
    {"code": "category_code", "title": "品类代码", "required": True, "type": "enum",
     "aliases": ["categorycode", "category", "品类代码", "所属品类", "品类", "产品分类", "商品分类"]},
    {"code": "lifecycle_status", "title": "生命周期", "required": False, "type": "enum",
     "aliases": ["lifecyclestatus", "lifecycle", "生命周期", "产品状态"]},
    {"code": "description", "title": "产品说明", "required": False, "type": "text",
     "aliases": ["description", "产品说明", "说明", "描述"]},
    {"code": "length", "title": "长度", "required": False, "type": "decimal",
     "aliases": ["length", "长度", "长"]},
    {"code": "width", "title": "宽度", "required": False, "type": "decimal",
     "aliases": ["width", "宽度", "宽"]},
    {"code": "height", "title": "高度", "required": False, "type": "decimal",
     "aliases": ["height", "高度", "高"]},
    {"code": "dimension_unit", "title": "尺寸单位", "required": False, "type": "enum",
     "aliases": ["dimensionunit", "尺寸单位", "长宽高单位"]},
    {"code": "weight", "title": "重量", "required": False, "type": "decimal",
     "aliases": ["weight", "重量", "净重"]},
    {"code": "weight_unit", "title": "重量单位", "required": False, "type": "enum",
     "aliases": ["weightunit", "重量单位"]},
    {"code": "moq", "title": "MOQ", "required": False, "type": "decimal",
     "aliases": ["moq", "最小起订量", "起订量"]},
    {"code": "factory_price", "title": "出厂价", "required": False, "type": "decimal",
     "aliases": ["factoryprice", "出厂价", "工厂价"]},
    {"code": "currency", "title": "币种", "required": False, "type": "enum",
     "aliases": ["currency", "币种", "货币"]},
    {"code": "spu_code", "title": "SPU编码", "required": False, "type": "text",
     "aliases": ["spucode", "spu", "spu编码"]},
    {"code": "bundle_code", "title": "套装编码", "required": False, "type": "text",
     "aliases": ["bundlecode", "套装编码", "套装"]},
    {"code": "bundle_quantity", "title": "套装数量", "required": False, "type": "decimal",
     "aliases": ["bundlequantity", "套装数量"]},
    {"code": "bom_code", "title": "BOM编码", "required": False, "type": "text",
     "aliases": ["bomcode", "bom", "bom编码"]},
    {"code": "bom_quantity", "title": "BOM用量", "required": False, "type": "decimal",
     "aliases": ["bomquantity", "bom用量", "组件用量"]},
]
FIELD_CODES = {item["code"] for item in FIELD_SCHEMA}
FIELD_BY_CODE = {item["code"]: item for item in FIELD_SCHEMA}
DICTIONARIES = {
    "category_code": ["sofa", "chair", "table", "bed", "storage", "other"],
    "lifecycle_status": ["concept", "sample", "active", "discontinued"],
    "dimension_unit": ["mm", "cm", "m", "in"],
    "weight_unit": ["g", "kg", "lb"],
    "currency": ["CNY", "USD", "EUR", "GBP"],
}
ENUM_VALUE_ALIASES = {
    "category_code": {
        "沙发": "sofa", "椅子": "chair", "椅类": "chair", "座椅": "chair",
        "桌子": "table", "桌类": "table", "床": "bed", "床类": "bed",
        "收纳": "storage", "收纳类": "storage", "其他": "other",
    },
    "lifecycle_status": {
        "概念": "concept", "打样": "sample", "样品": "sample", "在售": "active",
        "启用": "active", "停售": "discontinued", "停产": "discontinued",
    },
    "dimension_unit": {
        "毫米": "mm", "厘米": "cm", "米": "m", "英寸": "in",
    },
    "weight_unit": {
        "克": "g", "千克": "kg", "公斤": "kg", "磅": "lb",
    },
    "currency": {
        "人民币": "CNY", "人民币元": "CNY", "美元": "USD",
        "欧元": "EUR", "英镑": "GBP",
    },
}
REPAIR_SUGGESTIONS = {
    "REQUIRED": "补充必填值后重新预检",
    "TEXT_TYPE": "将单元格格式改为文本，避免编码前导零丢失",
    "TEXT_TOO_LONG": "缩短文本至字段允许长度",
    "NUMBER_INVALID": "填写可解析的非负数字",
    "ENUM_UNKNOWN": "使用 Dictionaries 工作表中的标准代码，或显式配置字典映射",
    "UNIT_REQUIRED": "填写与数值匹配的标准单位",
    "UNIT_WITHOUT_VALUE": "补充对应数值或清空单位",
    "SKU_DUPLICATE_FILE": "保留一行，或改为不同 SKU",
    "SKU_EXISTS_SKIPPED": "无需处理；如需覆盖已有主档，请选择“新增并更新”",
    "FORMULA_CELL": "粘贴公式计算结果为静态值后重新上传",
    "ALIAS_DUPLICATE": "来源编码和产品 SKU 必须保持一一对应",
    "ALIAS_TARGET_MISSING": "先创建产品主档，或修正目标产品 SKU",
}


def _json_default(value: Any) -> Any:
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return str(value)
    return str(value)


def _stable_json(value: Any) -> str:
    return json.dumps(
        value, ensure_ascii=False, sort_keys=True, separators=(",", ":"),
        default=_json_default,
    )


def _sha256(value: bytes | str) -> str:
    payload = value.encode("utf-8") if isinstance(value, str) else value
    return hashlib.sha256(payload).hexdigest()


def _header_key(value: Any) -> str:
    return re.sub(r"[\s_\-（）()]+", "", str(value or "").strip().lower())


def _compare_key(value: Any) -> str:
    return str(value or "").strip().upper()


def _issue(code: str, field: str, message: str) -> dict[str, str]:
    return {
        "code": code,
        "field": field,
        "message": message,
        "suggestion": REPAIR_SUGGESTIONS.get(code, "修正后重新预检"),
    }


def template_schema() -> dict[str, Any]:
    return {
        "template_version": TEMPLATE_VERSION,
        "fields": [{key: value for key, value in field.items() if key != "aliases"}
                   for field in FIELD_SCHEMA],
        "dictionaries": DICTIONARIES,
        "sheets": ["填写说明", "Product_Master", "SKU_Alias", "Dictionaries", "Examples"],
    }


def template_schema_sha256() -> str:
    return _sha256(_stable_json(template_schema()))


def build_template() -> tuple[bytes, str, str]:
    workbook = Workbook()
    guide = workbook.active
    guide.title = "填写说明"
    schema_sha = template_schema_sha256()
    guide_rows = [
        ["FurniScope 产品主档导入模板", ""],
        ["template_version", TEMPLATE_VERSION],
        ["schema_sha256", schema_sha],
        ["正式导入区", "Product_Master；不得放示例数据"],
        ["必填字段", "SKU、产品名称、品类代码"],
        ["SKU 规则", "按去空格后的大小写无关比较键保持租户内唯一；原始文本原样保存"],
        ["更新规则", "默认仅新增；批量更新必须在预检时显式选择 upsert"],
        ["尺寸规则", "长度、宽度或高度存在时，尺寸单位必填"],
        ["枚举规则", "只接受 Dictionaries 中的标准代码，不猜测未知值"],
        ["公式规则", "正式导入区禁止公式单元格，请粘贴静态值"],
    ]
    for row in guide_rows:
        guide.append(row)

    master = workbook.create_sheet("Product_Master")
    master.append([field["title"] for field in FIELD_SCHEMA])
    alias = workbook.create_sheet("SKU_Alias")
    alias.append(["来源系统", "来源SKU", "产品SKU"])

    dictionaries = workbook.create_sheet("Dictionaries")
    dictionary_codes = list(DICTIONARIES)
    dictionaries.append(dictionary_codes)
    max_values = max(map(len, DICTIONARIES.values()))
    for index in range(max_values):
        dictionaries.append([
            DICTIONARIES[code][index] if index < len(DICTIONARIES[code]) else None
            for code in dictionary_codes
        ])

    examples = workbook.create_sheet("Examples")
    examples.append([field["title"] for field in FIELD_SCHEMA])
    examples.append([
        "HF-EXAMPLE-001", "示例模块沙发", "sofa", "sample", "仅供参考，不参与正式导入",
        225, 95, 86, "cm", 68, "kg", 20, 299, "USD",
        "SPU-EXAMPLE", None, None, "BOM-EXAMPLE", 1,
    ])
    examples.append([
        "HF-EXAMPLE-002", "示例茶几", "table", "concept", None,
        120, 60, 42, "cm", 18, "kg", 50, 89, "USD",
        "SPU-TABLE-EXAMPLE", "SET-LIVING-EXAMPLE", 1, None, None,
    ])

    header_fill = PatternFill("solid", fgColor="0B918B")
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
            cell.alignment = Alignment(vertical="center")
        for column in sheet.columns:
            letter = column[0].column_letter
            sheet.column_dimensions[letter].width = min(
                34, max(12, max(len(str(cell.value or "")) for cell in column) + 2)
            )

    for code, values in DICTIONARIES.items():
        column = next(
            index for index, field in enumerate(FIELD_SCHEMA, start=1)
            if field["code"] == code
        )
        dictionary_column = dictionary_codes.index(code) + 1
        letter = dictionaries.cell(row=1, column=dictionary_column).column_letter
        validation = DataValidation(
            type="list",
            formula1=f"'Dictionaries'!${letter}$2:${letter}${len(values) + 1}",
            allow_blank=not FIELD_BY_CODE[code]["required"],
        )
        master.add_data_validation(validation)
        validation.add(f"{master.cell(row=2, column=column).column_letter}2:"
                       f"{master.cell(row=2, column=column).column_letter}20001")

    buffer = BytesIO()
    workbook.save(buffer)
    content = buffer.getvalue()
    return content, _sha256(content), schema_sha


def suggest_mapping(headers: list[str]) -> dict[str, str]:
    by_key: dict[str, list[str]] = {}
    for header in headers:
        by_key.setdefault(_header_key(header), []).append(header)
    output: dict[str, str] = {}
    for field in FIELD_SCHEMA:
        matches = {
            candidate
            for alias in field["aliases"]
            for candidate in by_key.get(_header_key(alias), [])
        }
        if len(matches) == 1:
            output[field["code"]] = matches.pop()
    return output


def _enum_target(code: str, value: Any) -> str | None:
    source = str(value or "").strip()
    if not source:
        return None
    canonical = {
        _header_key(item): item for item in DICTIONARIES.get(code, [])
    }
    aliases = {
        _header_key(item): target
        for item, target in ENUM_VALUE_ALIASES.get(code, {}).items()
    }
    return canonical.get(_header_key(source)) or aliases.get(_header_key(source))


def _header_similarity(field: dict[str, Any], header: str) -> int:
    key = _header_key(header)
    alias_keys = [_header_key(alias) for alias in field["aliases"]]
    if key in alias_keys:
        return 100
    if any(
        len(alias) >= 2 and len(key) >= 2 and (alias in key or key in alias)
        for alias in alias_keys
    ):
        return 72
    return 0


def _column_type_score(code: str, values: list[Any]) -> tuple[int, bool]:
    nonempty = [value for value in values if value not in (None, "")]
    if not nonempty:
        return 0, False
    field = FIELD_BY_CODE[code]
    if field["type"] == "decimal":
        valid = 0
        for value in nonempty:
            try:
                parsed = Decimal(str(value).replace(",", "").strip())
                valid += int(parsed.is_finite() and parsed >= 0)
            except (InvalidOperation, ValueError):
                pass
        ratio = valid / len(nonempty)
        return round(ratio * 22), ratio < 0.8
    if code in DICTIONARIES:
        valid = sum(_enum_target(code, value) is not None for value in nonempty)
        ratio = valid / len(nonempty)
        return round(ratio * 22), False
    text_values = [value for value in nonempty if isinstance(value, str)]
    if code == "sku":
        text_ratio = len(text_values) / len(nonempty)
        code_ratio = sum(
            bool(re.search(r"[A-Za-z0-9]", value)) and len(value.strip()) <= 100
            for value in text_values
        ) / len(nonempty)
        unique_ratio = len({value.strip().upper() for value in text_values}) / len(nonempty)
        return round((text_ratio * 8) + (code_ratio * 8) + (unique_ratio * 6)), text_ratio < 0.8
    text_ratio = len(text_values) / len(nonempty)
    return round(text_ratio * 14), False


def _row_headers(sheet, row_number: int) -> list[str]:
    return [
        str(sheet.cell(row=row_number, column=column).value or "").strip()
        for column in range(1, sheet.max_column + 1)
    ]


def _detect_table(workbook, *, requested_sheet: str | None = None) -> tuple[str, int]:
    excluded = {"填写说明", "SKU_Alias", "Dictionaries", "Examples"}
    sheet_names = [requested_sheet] if requested_sheet else list(workbook.sheetnames)
    best: tuple[int, str, int] | None = None
    for name in sheet_names:
        if name not in workbook.sheetnames:
            continue
        sheet = workbook[name]
        for row_number in range(1, min(100, sheet.max_row) + 1):
            headers = [header for header in _row_headers(sheet, row_number) if header]
            if len(headers) < 2:
                continue
            mapping = suggest_mapping(headers)
            required = sum(code in mapping for code in ("sku", "name", "category_code"))
            score = required * 100 + len(mapping) * 10 + min(len(headers), 30)
            if name == "Product_Master":
                score += 500
            elif name in excluded:
                score -= 400
            score -= row_number
            candidate = (score, name, row_number)
            if best is None or candidate > best:
                best = candidate
    if best is not None:
        return best[1], best[2]
    fallback = next(
        (name for name in workbook.sheetnames if name not in excluded),
        workbook.sheetnames[0],
    )
    return fallback, 1


def _inspect_loaded_workbook(
    workbook, *, filename: str, requested_sheet: str | None = None
) -> dict[str, Any]:
    sheet_name, header_row = _detect_table(
        workbook, requested_sheet=requested_sheet
    )
    sheet = workbook[sheet_name]
    headers = [header for header in _row_headers(sheet, header_row) if header]
    header_columns = {
        header: index
        for index, header in enumerate(_row_headers(sheet, header_row), start=1)
        if header
    }
    samples = {
        header: [
            sheet.cell(row=row, column=column).value
            for row in range(header_row + 1, min(sheet.max_row, header_row + 80) + 1)
        ]
        for header, column in header_columns.items()
    }
    mapping: dict[str, str] = {}
    detected_fields: list[dict[str, Any]] = []
    mapping_issues: list[dict[str, Any]] = []
    used_headers: set[str] = set()

    for field in FIELD_SCHEMA:
        candidates = []
        for header in headers:
            header_score = _header_similarity(field, header)
            type_score, type_conflict = _column_type_score(
                field["code"], samples.get(header, [])
            )
            if header_score or (field["required"] and type_score >= 18):
                candidates.append({
                    "header": header,
                    "score": min(100, header_score + type_score),
                    "header_score": header_score,
                    "type_conflict": type_conflict,
                })
        candidates.sort(key=lambda item: (-item["score"], item["header"]))
        exact = [item for item in candidates if item["header_score"] == 100]
        chosen = None
        issue_kind = None
        if len(exact) == 1:
            chosen = exact[0]
        elif len(exact) > 1:
            issue_kind = "ambiguous"
        elif candidates and candidates[0]["score"] >= 88 and (
            len(candidates) == 1 or candidates[0]["score"] - candidates[1]["score"] >= 12
        ):
            chosen = candidates[0]
        elif len(candidates) > 1 and candidates[0]["score"] >= 70:
            issue_kind = "ambiguous"

        if chosen and chosen["header"] not in used_headers:
            mapping[field["code"]] = chosen["header"]
            used_headers.add(chosen["header"])
            detected_fields.append({
                "field_code": field["code"],
                "title": field["title"],
                "source_header": chosen["header"],
                "confidence": chosen["score"],
            })
            if chosen["type_conflict"]:
                mapping_issues.append({
                    "field_code": field["code"],
                    "title": field["title"],
                    "kind": "type_conflict",
                    "required": field["required"],
                    "current_header": chosen["header"],
                    "candidates": [item["header"] for item in candidates[:5]],
                    "message": f"{chosen['header']} 列的数据类型不一致，请确认字段来源",
                })
        elif field["required"] or issue_kind:
            kind = issue_kind or "missing_required"
            issue_candidates = (
                exact
                if kind == "ambiguous" and len(exact) > 1
                else [item for item in candidates if item["header_score"] > 0]
            )
            mapping_issues.append({
                "field_code": field["code"],
                "title": field["title"],
                "kind": kind,
                "required": field["required"],
                "current_header": None,
                "candidates": [item["header"] for item in issue_candidates[:5]] or headers,
                "message": (
                    f"多个列都可能是{field['title']}，请选择正确来源"
                    if kind == "ambiguous"
                    else f"未识别到必填字段{field['title']}，文件中真实缺失时请补充后重新上传"
                ),
            })

    dictionary_mapping: dict[str, dict[str, str]] = {}
    unit_mapping: dict[str, str] = {}
    value_issues: list[dict[str, Any]] = []
    for code, allowed in DICTIONARIES.items():
        header = mapping.get(code)
        if not header:
            continue
        source_values = sorted({
            str(value).strip()
            for value in samples.get(header, [])
            if value not in (None, "")
        })
        unresolved = []
        for source in source_values:
            target = _enum_target(code, source)
            if target is None:
                unresolved.append(source)
            elif target != source:
                if code in {"dimension_unit", "weight_unit"}:
                    unit_mapping[source] = target
                else:
                    dictionary_mapping.setdefault(code, {})[source] = target
        if unresolved:
            value_issues.append({
                "field_code": code,
                "title": FIELD_BY_CODE[code]["title"],
                "source_values": unresolved[:20],
                "allowed_values": allowed,
                "message": f"{FIELD_BY_CODE[code]['title']}存在未识别值，请确认标准代码",
            })

    total_rows = sum(
        any(
            sheet.cell(row=row, column=column).value not in (None, "")
            for column in header_columns.values()
        )
        for row in range(header_row + 1, sheet.max_row + 1)
    )
    header_keys = [_header_key(header) for header in headers]
    warnings = []
    if len(header_keys) != len(set(header_keys)):
        warnings.append("表头包含重复列名，请修改后重新上传")
    return {
        "original_filename": filename,
        "sheet_name": sheet_name,
        "available_sheets": list(workbook.sheetnames),
        "header_row": header_row,
        "headers": headers,
        "total_rows": total_rows,
        "field_mapping": mapping,
        "detected_fields": detected_fields,
        "mapping_issues": mapping_issues,
        "value_issues": value_issues,
        "unit_mapping": unit_mapping,
        "dictionary_mapping": dictionary_mapping,
        "ignored_columns": [header for header in headers if header not in mapping.values()],
        "warnings": warnings,
    }


def _decimal(value: Any, field: str, errors: list[dict[str, str]]) -> str | None:
    if value is None or value == "":
        return None
    try:
        parsed = Decimal(str(value).replace(",", "").strip())
        if not parsed.is_finite() or parsed < 0:
            raise InvalidOperation
    except (InvalidOperation, ValueError):
        errors.append(_issue("NUMBER_INVALID", field, f"{FIELD_BY_CODE[field]['title']}必须是非负数字"))
        return None
    return format(parsed.normalize(), "f")


def _normalize_values(
    values: dict[str, Any], *, dictionary_mapping: dict[str, dict[str, str]] | None = None,
    unit_mapping: dict[str, str] | None = None, sku_from_cell: Any = None,
) -> tuple[dict[str, Any], list[dict[str, str]], list[dict[str, str]]]:
    dictionary_mapping = dictionary_mapping or {}
    unit_mapping = unit_mapping or {}
    normalized: dict[str, Any] = {}
    errors: list[dict[str, str]] = []
    warnings: list[dict[str, str]] = []

    # Schema order keeps validation output and the preview SHA stable across processes.
    for field in FIELD_SCHEMA:
        code = field["code"]
        value = values.get(code)
        if isinstance(value, str):
            value = value.strip()
        if value == "":
            value = None
        if code in {"sku", "name"}:
            if value is None:
                errors.append(_issue("REQUIRED", code, f"{FIELD_BY_CODE[code]['title']}不能为空"))
                continue
            original = sku_from_cell if code == "sku" and sku_from_cell is not None else value
            if not isinstance(original, str):
                errors.append(_issue("TEXT_TYPE", code, "SKU 必须使用文本单元格，不能使用数值编码"))
                continue
            value = str(value).strip()
            maximum = int(FIELD_BY_CODE[code]["max_length"])
            if len(value) > maximum:
                errors.append(_issue("TEXT_TOO_LONG", code, f"最多允许 {maximum} 个字符"))
                continue
            normalized[code] = value
        elif code in {"description", "spu_code", "bundle_code", "bom_code"}:
            if value is not None:
                normalized[code] = str(value).strip()
        elif FIELD_BY_CODE[code]["type"] == "decimal":
            parsed = _decimal(value, code, errors)
            if parsed is not None:
                normalized[code] = parsed
        elif code in DICTIONARIES:
            if value is None:
                if code == "category_code":
                    errors.append(_issue("REQUIRED", code, "品类代码不能为空"))
                continue
            source = str(value).strip()
            mapped = dictionary_mapping.get(code, {}).get(source, source)
            if code in {"dimension_unit", "weight_unit"}:
                mapped = unit_mapping.get(source, mapped)
            if mapped == source:
                mapped = _enum_target(code, source) or source
            if mapped not in DICTIONARIES[code]:
                errors.append(_issue(
                    "ENUM_UNKNOWN", code,
                    f"{source} 不在 {code} 的受控词表中",
                ))
                continue
            normalized[code] = mapped

    normalized.setdefault("lifecycle_status", "active")
    dimensions = [normalized.get(key) for key in ("length", "width", "height")]
    if any(value is not None for value in dimensions) and not normalized.get("dimension_unit"):
        errors.append(_issue("UNIT_REQUIRED", "dimension_unit", "填写尺寸值时必须填写尺寸单位"))
    if normalized.get("dimension_unit") and not any(value is not None for value in dimensions):
        warnings.append(_issue("UNIT_WITHOUT_VALUE", "dimension_unit", "尺寸单位存在，但未填写尺寸值"))
    if normalized.get("weight") is not None and not normalized.get("weight_unit"):
        errors.append(_issue("UNIT_REQUIRED", "weight_unit", "填写重量时必须填写重量单位"))
    if normalized.get("weight_unit") and normalized.get("weight") is None:
        warnings.append(_issue("UNIT_WITHOUT_VALUE", "weight_unit", "重量单位存在，但未填写重量"))
    if normalized.get("factory_price") is not None and not normalized.get("currency"):
        errors.append(_issue("UNIT_REQUIRED", "currency", "填写出厂价时必须填写币种"))
    if normalized.get("bundle_code") and normalized.get("bundle_quantity") is None:
        normalized["bundle_quantity"] = "1"
    if normalized.get("bom_code") and normalized.get("bom_quantity") is None:
        normalized["bom_quantity"] = "1"
    if normalized.get("category_code") not in {None, "sofa"}:
        warnings.append({
            "code": "ANALYSIS_CATEGORY_UNVERIFIED",
            "field": "category_code",
            "message": "该品类可进入产品主档，但当前多模态画像分析仅完成 sofa 验证",
            "suggestion": "导入后先人工维护画像，不要直接依赖自动分析结论",
        })
    return normalized, errors, warnings


class ProductImportService:
    def __init__(self) -> None:
        self.products = ProductRepository()
        self.audit = AuditService()

    async def _job_row(
        self, session: AsyncSession, *, tenant_id: int, job_uuid: str, lock: bool = False
    ) -> dict[str, Any]:
        row = (await session.execute(text("""
            SELECT id,job_uuid::text AS job_uuid,original_filename,source_sha256,
                   template_version,sheet_name,available_sheets,header_row,field_mapping,
                   unit_mapping,dictionary_mapping,import_mode,status,total_rows,valid_rows,
                   error_rows,warning_rows,imported_rows,preview_sha256,schema_snapshot,
                   error_message,created_at,updated_at,committed_at
              FROM product_import_jobs
             WHERE tenant_id=:tenant_id AND job_uuid=CAST(:job_uuid AS uuid)
        """ + (" FOR UPDATE" if lock else "")),
            {"tenant_id": tenant_id, "job_uuid": job_uuid})).mappings().one_or_none()
        if row is None:
            raise BusinessError("PRODUCT_IMPORT_NOT_FOUND", "导入任务不存在或不可访问", status_code=404)
        return dict(row)

    async def job(
        self, session: AsyncSession, *, tenant_id: int, job_uuid: str,
        include_rows: bool = True, row_limit: int = 500,
    ) -> dict[str, Any]:
        job = await self._job_row(
            session, tenant_id=tenant_id, job_uuid=job_uuid
        )
        job_id = job.pop("id")
        rows: list[dict[str, Any]] = []
        if include_rows:
            result = await session.execute(text("""
                SELECT source_row_number,source_values,normalized_values,
                       validation_errors,validation_warnings,planned_action,
                       is_user_edited,imported_product_id
                  FROM product_import_rows
                 WHERE tenant_id=:tenant_id AND job_id=:job_id
                 ORDER BY source_row_number LIMIT :row_limit
            """), {"tenant_id": tenant_id, "job_id": job_id, "row_limit": row_limit})
            rows = [dict(row) for row in result.mappings().all()]
        return job | {"rows": rows}

    async def recent(
        self, session: AsyncSession, *, tenant_id: int, limit: int = 10
    ) -> list[dict[str, Any]]:
        rows = (await session.execute(text("""
            SELECT job_uuid::text AS job_uuid,original_filename,source_sha256,
                   template_version,sheet_name,available_sheets,header_row,field_mapping,
                   unit_mapping,dictionary_mapping,import_mode,status,total_rows,valid_rows,
                   error_rows,warning_rows,imported_rows,preview_sha256,schema_snapshot,
                   error_message,created_at,updated_at,committed_at
              FROM product_import_jobs
             WHERE tenant_id=:tenant_id
             ORDER BY created_at DESC LIMIT :limit
        """), {"tenant_id": tenant_id, "limit": limit})).mappings().all()
        return [dict(row) | {"rows": []} for row in rows]

    @staticmethod
    def _load_workbook(
        *, filename: str, mime_type: str, content: bytes
    ) -> tuple[str, Any]:
        safe_name = Path(filename or "").name
        if not safe_name.lower().endswith(".xlsx"):
            raise BusinessError("PRODUCT_IMPORT_FILE_TYPE", "仅支持 .xlsx 文件", status_code=422)
        if mime_type not in XLSX_MIME_TYPES:
            raise BusinessError("PRODUCT_IMPORT_MIME_TYPE", "文件 MIME 类型与 XLSX 不匹配", status_code=422)
        if not content or len(content) > MAX_FILE_SIZE:
            raise BusinessError("PRODUCT_IMPORT_FILE_SIZE", "文件必须大于 0 且不超过 20 MB", status_code=422)
        try:
            workbook = load_workbook(BytesIO(content), data_only=False, read_only=False)
        except Exception as exc:
            raise BusinessError(
                "PRODUCT_IMPORT_WORKBOOK_INVALID", "无法解析 XLSX 文件", status_code=422
            ) from exc
        if not workbook.sheetnames or len(workbook.sheetnames) > MAX_SHEETS:
            raise BusinessError(
                "PRODUCT_IMPORT_SHEET_COUNT", "工作表数量必须在 1 到 20 之间", status_code=422
            )
        return safe_name, workbook

    def inspect(
        self, *, filename: str, mime_type: str, content: bytes
    ) -> dict[str, Any]:
        safe_name, workbook = self._load_workbook(
            filename=filename, mime_type=mime_type, content=content
        )
        inspection = _inspect_loaded_workbook(workbook, filename=safe_name)
        sheet = workbook[inspection["sheet_name"]]
        if sheet.max_row > MAX_ROWS + inspection["header_row"]:
            raise BusinessError(
                "PRODUCT_IMPORT_ROW_LIMIT", "单次最多导入 20,000 行", status_code=422
            )
        workbook.close()
        return inspection

    async def preflight(
        self, session: AsyncSession, *, tenant_id: int, user_id: int,
        idempotency_key: str, filename: str, mime_type: str, content: bytes,
        sheet_name: str | None, header_row: int | None, field_mapping: dict[str, str],
        unit_mapping: dict[str, str], dictionary_mapping: dict[str, dict[str, str]],
        import_mode: str, request_id: str,
    ) -> dict[str, Any]:
        if import_mode not in {"create_only", "upsert"}:
            raise BusinessError("PRODUCT_IMPORT_MODE", "导入模式无效", status_code=422)

        safe_name, workbook = self._load_workbook(
            filename=filename, mime_type=mime_type, content=content
        )
        if sheet_name is not None and sheet_name not in workbook.sheetnames:
            raise BusinessError("PRODUCT_IMPORT_SHEET_NOT_FOUND", "选择的工作表不存在", status_code=422)
        inspection = _inspect_loaded_workbook(
            workbook, filename=safe_name, requested_sheet=sheet_name
        )
        selected_name = sheet_name or inspection["sheet_name"]
        header_row = header_row or inspection["header_row"]
        if not 1 <= header_row <= 100:
            raise BusinessError("PRODUCT_IMPORT_HEADER_ROW", "表头行必须在 1 到 100 之间", status_code=422)
        auto_unit_mapping = inspection["unit_mapping"]
        unit_mapping = auto_unit_mapping | unit_mapping
        auto_dictionary_mapping = inspection["dictionary_mapping"]
        dictionary_mapping = {
            code: auto_dictionary_mapping.get(code, {}) | dictionary_mapping.get(code, {})
            for code in set(auto_dictionary_mapping) | set(dictionary_mapping)
        }
        detected_mapping = inspection["field_mapping"]
        field_mapping = detected_mapping | field_mapping
        source_sha = _sha256(content)
        config_snapshot = {
            "sheet_name": selected_name, "header_row": header_row,
            "field_mapping": field_mapping, "unit_mapping": unit_mapping,
            "dictionary_mapping": dictionary_mapping, "import_mode": import_mode,
        }
        existing = (await session.execute(text("""
            SELECT job_uuid::text AS job_uuid,source_sha256,schema_snapshot
              FROM product_import_jobs
             WHERE tenant_id=:tenant_id AND idempotency_key=:idempotency_key
        """), {"tenant_id": tenant_id,
                 "idempotency_key": idempotency_key})).mappings().one_or_none()
        if existing:
            if (existing["source_sha256"] != source_sha
                    or existing["schema_snapshot"].get("request_config") != config_snapshot):
                raise BusinessError("IDEMPOTENCY_CONFLICT", "幂等键已用于不同的导入文件或配置", status_code=409)
            return await self.job(
                session, tenant_id=tenant_id, job_uuid=str(existing["job_uuid"])
            )

        sheet = workbook[selected_name]
        if sheet.max_row > MAX_ROWS + header_row:
            raise BusinessError("PRODUCT_IMPORT_ROW_LIMIT", "单次最多导入 20,000 行", status_code=422)

        header_values = [
            str(sheet.cell(row=header_row, column=column).value or "").strip()
            for column in range(1, sheet.max_column + 1)
        ]
        nonempty_headers = [header for header in header_values if header]
        header_keys = [_header_key(header) for header in nonempty_headers]
        file_errors: list[dict[str, str]] = []
        if len(header_keys) != len(set(header_keys)):
            file_errors.append(_issue("HEADER_DUPLICATE", "_file", "表头包含重复列名"))
        suggested = suggest_mapping(nonempty_headers) | detected_mapping
        mapping = suggested | field_mapping
        unknown_fields = set(mapping) - FIELD_CODES
        unknown_headers = set(mapping.values()) - set(nonempty_headers)
        if unknown_fields or unknown_headers:
            raise BusinessError("PRODUCT_IMPORT_MAPPING_INVALID", "字段映射包含未知字段或列名", status_code=422)
        missing_required = [
            FIELD_BY_CODE[code]["title"] for code in ("sku", "name", "category_code")
            if code not in mapping
        ]
        if missing_required:
            file_errors.append(_issue(
                "MAPPING_REQUIRED", "_file",
                f"尚未映射必填字段：{'、'.join(missing_required)}",
            ))

        header_columns = {
            header: index for index, header in enumerate(header_values, start=1) if header
        }
        parsed_rows: list[dict[str, Any]] = []
        formula_coordinates: list[str] = []
        for row_number in range(header_row + 1, sheet.max_row + 1):
            source = {
                header: sheet.cell(row=row_number, column=column).value
                for header, column in header_columns.items()
            }
            if not any(value is not None and value != "" for value in source.values()):
                continue
            values = {
                code: source.get(header) for code, header in mapping.items()
            }
            formula_fields = [
                code for code, header in mapping.items()
                if sheet.cell(row=row_number, column=header_columns[header]).data_type == "f"
            ]
            formula_coordinates.extend(
                sheet.cell(row=row_number, column=header_columns[mapping[code]]).coordinate
                for code in formula_fields
            )
            normalized, errors, warnings = _normalize_values(
                values, dictionary_mapping=dictionary_mapping,
                unit_mapping=unit_mapping, sku_from_cell=values.get("sku"),
            )
            errors.extend(
                _issue("FORMULA_CELL", code, f"{FIELD_BY_CODE[code]['title']}不能使用公式")
                for code in formula_fields
            )
            parsed_rows.append({
                "source_row_number": row_number,
                "source_values": json.loads(_stable_json(source)),
                "normalized_values": normalized,
                "validation_errors": errors,
                "validation_warnings": warnings,
                "planned_action": "invalid",
                "is_user_edited": False,
            })
        if not parsed_rows:
            file_errors.append(_issue("NO_DATA_ROWS", "_file", "工作表中没有可导入的数据行"))

        keys = [_compare_key(row["normalized_values"].get("sku")) for row in parsed_rows]
        existing_products: dict[str, int] = {}
        if any(keys):
            result = await session.execute(text("""
                SELECT sku_compare_key,id FROM products
                 WHERE tenant_id=:tenant_id AND deleted_at IS NULL
                   AND sku_compare_key=ANY(:keys)
            """), {"tenant_id": tenant_id, "keys": list({key for key in keys if key})})
            existing_products = {
                row["sku_compare_key"]: int(row["id"]) for row in result.mappings().all()
            }
        all_existing_product_keys = set((await session.execute(text("""
            SELECT sku_compare_key FROM products
             WHERE tenant_id=:tenant_id AND deleted_at IS NULL
        """), {"tenant_id": tenant_id})).scalars().all())
        self._apply_cross_row_rules(parsed_rows, existing_products, import_mode)
        aliases, alias_errors = await self._parse_aliases(
            session, workbook=workbook, tenant_id=tenant_id,
            import_product_keys={key for key in keys if key},
            existing_product_keys=all_existing_product_keys,
        )
        schema_snapshot = {
            "request_config": config_snapshot,
            "schema_sha256": template_schema_sha256(),
            "field_schema": template_schema()["fields"],
            "suggested_mapping": suggested,
            "file_errors": file_errors,
            "formula_cells": formula_coordinates[:200],
            "aliases": aliases,
            "alias_errors": alias_errors,
            "preview_row_limit": 500,
        }
        preview_sha = self._preview_sha(parsed_rows, aliases)
        error_rows = sum(bool(row["validation_errors"]) for row in parsed_rows)
        valid_rows = len(parsed_rows) - error_rows
        warning_rows = sum(bool(row["validation_warnings"]) for row in parsed_rows)
        status = "ready" if not file_errors and not alias_errors and not error_rows else "blocked"
        template_version = None
        if "填写说明" in workbook.sheetnames:
            guide = workbook["填写说明"]
            for row in guide.iter_rows(min_row=1, max_col=2, values_only=True):
                if row[0] == "template_version":
                    template_version = str(row[1] or "") or None
                    break

        job_id, job_uuid = (await session.execute(text("""
            INSERT INTO product_import_jobs(
              tenant_id,created_by,idempotency_key,original_filename,mime_type,size_bytes,
              source_sha256,template_version,sheet_name,available_sheets,header_row,
              field_mapping,unit_mapping,dictionary_mapping,import_mode,status,total_rows,
              valid_rows,error_rows,warning_rows,preview_sha256,schema_snapshot
            ) VALUES(
              :tenant_id,:user_id,:idempotency_key,:filename,:mime_type,:size_bytes,
              :source_sha,:template_version,:sheet_name,CAST(:available_sheets AS jsonb),
              :header_row,CAST(:field_mapping AS jsonb),CAST(:unit_mapping AS jsonb),
              CAST(:dictionary_mapping AS jsonb),:import_mode,:status,:total_rows,
              :valid_rows,:error_rows,:warning_rows,:preview_sha,CAST(:schema_snapshot AS jsonb)
            )
            RETURNING id,job_uuid::text
        """), {
            "tenant_id": tenant_id, "user_id": user_id,
            "idempotency_key": idempotency_key, "filename": safe_name,
            "mime_type": mime_type, "size_bytes": len(content), "source_sha": source_sha,
            "template_version": template_version, "sheet_name": selected_name,
            "available_sheets": _stable_json(workbook.sheetnames),
            "header_row": header_row, "field_mapping": _stable_json(mapping),
            "unit_mapping": _stable_json(unit_mapping),
            "dictionary_mapping": _stable_json(dictionary_mapping),
            "import_mode": import_mode, "status": status,
            "total_rows": len(parsed_rows), "valid_rows": valid_rows,
            "error_rows": error_rows, "warning_rows": warning_rows,
            "preview_sha": preview_sha, "schema_snapshot": _stable_json(schema_snapshot),
        })).one()
        insert_row = text("""
            INSERT INTO product_import_rows(
              tenant_id,job_id,source_row_number,source_values,normalized_values,
              validation_errors,validation_warnings,planned_action,is_user_edited
            ) VALUES(
              :tenant_id,:job_id,:source_row_number,CAST(:source_values AS jsonb),
              CAST(:normalized_values AS jsonb),CAST(:validation_errors AS jsonb),
              CAST(:validation_warnings AS jsonb),:planned_action,:is_user_edited
            )
        """)
        row_parameters = [{
                "tenant_id": tenant_id, "job_id": job_id,
                **{key: (_stable_json(value) if key in {
                    "source_values", "normalized_values", "validation_errors",
                    "validation_warnings",
                } else value) for key, value in row.items()},
            } for row in parsed_rows]
        for offset in range(0, len(row_parameters), 1000):
            await session.execute(insert_row, row_parameters[offset:offset + 1000])
        await self.audit.record(
            session, tenant_id=tenant_id, actor_user_id=user_id,
            action_code="product.import.preflight", resource_type="product_import_job",
            resource_id=int(job_id), request_id=request_id,
            after={
                "job_uuid": str(job_uuid), "source_sha256": source_sha,
                "status": status, "total_rows": len(parsed_rows),
                "valid_rows": valid_rows, "error_rows": error_rows,
            },
        )
        return await self.job(
            session, tenant_id=tenant_id, job_uuid=str(job_uuid)
        )

    @staticmethod
    def _apply_cross_row_rules(
        rows: list[dict[str, Any]], existing_products: dict[str, int], import_mode: str
    ) -> None:
        by_key: dict[str, list[dict[str, Any]]] = {}
        for row in rows:
            key = _compare_key(row["normalized_values"].get("sku"))
            if key:
                by_key.setdefault(key, []).append(row)
        for duplicated in by_key.values():
            if len(duplicated) > 1:
                for row in duplicated:
                    row["validation_errors"].append(_issue(
                        "SKU_DUPLICATE_FILE", "sku", "文件内存在大小写或空格归一化后的重复 SKU"
                    ))
        for row in rows:
            key = _compare_key(row["normalized_values"].get("sku"))
            product_id = existing_products.get(key)
            if row["validation_errors"]:
                row["planned_action"] = "invalid"
            elif product_id and import_mode == "create_only":
                row["validation_warnings"].append(_issue(
                    "SKU_EXISTS_SKIPPED", "sku", "SKU 已存在，将按仅新增策略跳过"
                ))
                row["planned_action"] = "skip"
            else:
                row["planned_action"] = "update" if product_id else "create"

    async def _parse_aliases(
        self, session: AsyncSession, *, workbook, tenant_id: int,
        import_product_keys: set[str], existing_product_keys: set[str],
    ) -> tuple[list[dict[str, str]], list[dict[str, str]]]:
        if "SKU_Alias" not in workbook.sheetnames:
            return [], []
        sheet = workbook["SKU_Alias"]
        headers = {
            _header_key(cell.value): cell.column
            for cell in sheet[1] if cell.value is not None and cell.value != ""
        }
        source_column = headers.get(_header_key("来源SKU"))
        target_column = headers.get(_header_key("产品SKU"))
        system_column = headers.get(_header_key("来源系统"))
        if not source_column or not target_column:
            return [], [_issue("ALIAS_HEADER", "_alias", "SKU_Alias 缺少来源SKU或产品SKU列")]
        aliases: list[dict[str, str]] = []
        errors: list[dict[str, str]] = []
        source_keys: set[str] = set()
        target_keys: set[str] = set()
        for row_number in range(2, sheet.max_row + 1):
            source = sheet.cell(row=row_number, column=source_column).value
            target = sheet.cell(row=row_number, column=target_column).value
            if (source is None or source == "") and (target is None or target == ""):
                continue
            source_text, target_text = str(source or "").strip(), str(target or "").strip()
            source_key, target_key = _compare_key(source_text), _compare_key(target_text)
            if not source_key or not target_key:
                errors.append(_issue("REQUIRED", "_alias", f"SKU_Alias 第 {row_number} 行编码不能为空"))
                continue
            if source_key in source_keys or target_key in target_keys:
                errors.append(_issue("ALIAS_DUPLICATE", "_alias", f"SKU_Alias 第 {row_number} 行不是一一对应"))
                continue
            if target_key not in import_product_keys | existing_product_keys:
                errors.append(_issue("ALIAS_TARGET_MISSING", "_alias", f"目标产品 {target_text} 不存在"))
                continue
            source_keys.add(source_key)
            target_keys.add(target_key)
            aliases.append({
                "source_system": str(
                    sheet.cell(row=row_number, column=system_column).value or "default"
                ).strip() if system_column else "default",
                "source_sku": source_text,
                "product_sku": target_text,
            })
        if aliases:
            existing = (await session.execute(text("""
                SELECT source_sku,product_sku FROM tenant_forecast_sku_aliases
                 WHERE tenant_id=:tenant_id AND source_context=:context
                   AND (
                     upper(btrim(source_sku))=ANY(:sources)
                     OR upper(btrim(product_sku))=ANY(:targets)
                   )
            """), {
                "tenant_id": tenant_id, "context": contract_context(tenant_id),
                "sources": list(source_keys), "targets": list(target_keys),
            })).mappings().all()
            requested = {
                (_compare_key(item["source_sku"]), _compare_key(item["product_sku"]))
                for item in aliases
            }
            for row in existing:
                if (_compare_key(row["source_sku"]), _compare_key(row["product_sku"])) not in requested:
                    errors.append(_issue(
                        "ALIAS_DUPLICATE", "_alias",
                        "来源 SKU 或产品 SKU 已被现有映射占用",
                    ))
                    break
        return aliases, errors

    @staticmethod
    def _preview_sha(rows: list[dict[str, Any]], aliases: list[dict[str, str]]) -> str:
        payload = [{
            "row": row["source_row_number"],
            "normalized": row["normalized_values"],
            "errors": row["validation_errors"],
            "warnings": row["validation_warnings"],
            "action": row["planned_action"],
        } for row in rows]
        return _sha256(_stable_json({"rows": payload, "aliases": aliases}))

    async def patch_row(
        self, session: AsyncSession, *, tenant_id: int, user_id: int,
        job_uuid: str, source_row_number: int, normalized_values: dict[str, Any],
        request_id: str,
    ) -> dict[str, Any]:
        job = await self._job_row(
            session, tenant_id=tenant_id, job_uuid=job_uuid, lock=True
        )
        if job["status"] not in {"ready", "blocked", "failed"}:
            raise BusinessError("PRODUCT_IMPORT_IMMUTABLE", "当前导入任务不可编辑", status_code=409)
        unknown = set(normalized_values) - FIELD_CODES
        if unknown:
            raise BusinessError("PRODUCT_IMPORT_ROW_INVALID", "行数据包含未知字段", status_code=422)
        normalized, errors, warnings = _normalize_values(normalized_values)
        updated = await session.execute(text("""
            UPDATE product_import_rows
               SET normalized_values=CAST(:normalized AS jsonb),
                   validation_errors=CAST(:errors AS jsonb),
                   validation_warnings=CAST(:warnings AS jsonb),
                   is_user_edited=TRUE
             WHERE tenant_id=:tenant_id AND job_id=:job_id
               AND source_row_number=:source_row_number
            RETURNING source_row_number
        """), {
            "tenant_id": tenant_id, "job_id": job["id"],
            "source_row_number": source_row_number,
            "normalized": _stable_json(normalized), "errors": _stable_json(errors),
            "warnings": _stable_json(warnings),
        })
        if updated.scalar_one_or_none() is None:
            raise BusinessError("PRODUCT_IMPORT_ROW_NOT_FOUND", "导入行不存在", status_code=404)
        rows = [dict(row) for row in (await session.execute(text("""
            SELECT source_row_number,source_values,normalized_values,validation_errors,
                   validation_warnings,planned_action,is_user_edited,imported_product_id
              FROM product_import_rows
             WHERE tenant_id=:tenant_id AND job_id=:job_id
             ORDER BY source_row_number
        """), {"tenant_id": tenant_id, "job_id": job["id"]})).mappings().all()]
        keys = {_compare_key(row["normalized_values"].get("sku")) for row in rows}
        existing = (await session.execute(text("""
            SELECT sku_compare_key,id FROM products
             WHERE tenant_id=:tenant_id AND deleted_at IS NULL
               AND sku_compare_key=ANY(:keys)
        """), {"tenant_id": tenant_id, "keys": list(key for key in keys if key)})).mappings().all()
        existing_products = {row["sku_compare_key"]: int(row["id"]) for row in existing}
        for row in rows:
            values, row_errors, row_warnings = _normalize_values(row["normalized_values"])
            row["normalized_values"] = values
            row["validation_errors"] = row_errors
            row["validation_warnings"] = row_warnings
        self._apply_cross_row_rules(rows, existing_products, job["import_mode"])
        for row in rows:
            await session.execute(text("""
                UPDATE product_import_rows SET
                  normalized_values=CAST(:normalized AS jsonb),
                  validation_errors=CAST(:errors AS jsonb),
                  validation_warnings=CAST(:warnings AS jsonb),
                  planned_action=:action
                 WHERE tenant_id=:tenant_id AND job_id=:job_id
                   AND source_row_number=:source_row_number
            """), {
                "tenant_id": tenant_id, "job_id": job["id"],
                "source_row_number": row["source_row_number"],
                "normalized": _stable_json(row["normalized_values"]),
                "errors": _stable_json(row["validation_errors"]),
                "warnings": _stable_json(row["validation_warnings"]),
                "action": row["planned_action"],
            })
        aliases = job["schema_snapshot"].get("aliases", [])
        preview_sha = self._preview_sha(rows, aliases)
        error_rows = sum(bool(row["validation_errors"]) for row in rows)
        warning_rows = sum(bool(row["validation_warnings"]) for row in rows)
        status = (
            "ready" if not error_rows
            and not job["schema_snapshot"].get("file_errors")
            and not job["schema_snapshot"].get("alias_errors")
            else "blocked"
        )
        await session.execute(text("""
            UPDATE product_import_jobs SET status=:status,valid_rows=:valid_rows,
                   error_rows=:error_rows,warning_rows=:warning_rows,
                   preview_sha256=:preview_sha,error_message=NULL
             WHERE id=:job_id AND tenant_id=:tenant_id
        """), {
            "status": status, "valid_rows": len(rows) - error_rows,
            "error_rows": error_rows, "warning_rows": warning_rows,
            "preview_sha": preview_sha, "job_id": job["id"], "tenant_id": tenant_id,
        })
        await self.audit.record(
            session, tenant_id=tenant_id, actor_user_id=user_id,
            action_code="product.import.row.update", resource_type="product_import_job",
            resource_id=job["id"], request_id=request_id,
            after={"job_uuid": job_uuid, "source_row_number": source_row_number,
                   "status": status, "preview_sha256": preview_sha},
        )
        return await self.job(session, tenant_id=tenant_id, job_uuid=job_uuid)

    async def commit(
        self, session: AsyncSession, *, tenant_id: int, user_id: int,
        job_uuid: str, preview_sha256: str, request_id: str,
    ) -> dict[str, Any]:
        await session.execute(text("SELECT pg_advisory_xact_lock(:lock_key)"),
                              {"lock_key": 7_320_000 + tenant_id})
        job = await self._job_row(
            session, tenant_id=tenant_id, job_uuid=job_uuid, lock=True
        )
        if job["status"] == "completed":
            return {
                "job_uuid": job_uuid, "status": "completed",
                "imported_rows": job["imported_rows"],
                "created_rows": int(job["schema_snapshot"].get("created_rows", 0)),
                "updated_rows": int(job["schema_snapshot"].get("updated_rows", 0)),
                "skipped_rows": int(job["schema_snapshot"].get("skipped_rows", 0)),
                "alias_rows": int(job["schema_snapshot"].get("alias_rows", 0)),
                "committed_at": job["committed_at"],
            }
        if job["status"] != "ready":
            raise BusinessError("PRODUCT_IMPORT_BLOCKED", "预检仍有阻断错误，不能提交", status_code=422)
        if job["preview_sha256"] != preview_sha256:
            raise BusinessError("PRODUCT_IMPORT_PREVIEW_STALE", "预览已变化，请重新确认", status_code=409)
        rows = [dict(row) for row in (await session.execute(text("""
            SELECT source_row_number,normalized_values,planned_action
              FROM product_import_rows
             WHERE tenant_id=:tenant_id AND job_id=:job_id
             ORDER BY source_row_number FOR UPDATE
        """), {"tenant_id": tenant_id, "job_id": job["id"]})).mappings().all()]
        if len(rows) != job["valid_rows"] or any(row["planned_action"] == "invalid" for row in rows):
            raise BusinessError("PRODUCT_IMPORT_PREVIEW_STALE", "预检计数已变化，请重新预检", status_code=409)
        await session.execute(text("""
            UPDATE product_import_jobs SET status='importing'
             WHERE id=:job_id AND tenant_id=:tenant_id
        """), {"job_id": job["id"], "tenant_id": tenant_id})

        created_rows = 0
        updated_rows = 0
        skipped_rows = 0
        for row in rows:
            values = row["normalized_values"]
            if row["planned_action"] == "skip":
                skipped_rows += 1
                continue
            if row["planned_action"] == "create":
                product_id = await session.scalar(text("""
                    INSERT INTO products(
                      tenant_id,sku,name,category_code,lifecycle_status,description,created_by
                    ) VALUES(
                      :tenant_id,:sku,:name,:category_code,:lifecycle_status,:description,:user_id
                    )
                    RETURNING id
                """), {
                    "tenant_id": tenant_id, "user_id": user_id,
                    "sku": values["sku"], "name": values["name"],
                    "category_code": values["category_code"],
                    "lifecycle_status": values["lifecycle_status"],
                    "description": values.get("description"),
                })
                created_rows += 1
            else:
                product_id = await session.scalar(text("""
                    UPDATE products SET name=:name,category_code=:category_code,
                           lifecycle_status=:lifecycle_status,
                           description=COALESCE(:description,description),
                           analysis_status=CASE
                             WHEN CAST(:lifecycle_status AS varchar)='discontinued'
                               THEN 'archived'
                             WHEN category_code IS DISTINCT FROM
                                  CAST(:category_code AS varchar) THEN 'draft'
                             ELSE analysis_status END
                     WHERE tenant_id=:tenant_id AND deleted_at IS NULL
                       AND sku_compare_key=:sku_key
                    RETURNING id
                """), {
                    "tenant_id": tenant_id, "sku_key": _compare_key(values["sku"]),
                    "name": values["name"], "category_code": values["category_code"],
                    "lifecycle_status": values["lifecycle_status"],
                    "description": values.get("description"),
                })
                if product_id is None:
                    raise BusinessError(
                        "PRODUCT_IMPORT_PREVIEW_STALE",
                        f"第 {row['source_row_number']} 行目标产品已变化",
                        status_code=409,
                    )
                updated_rows += 1
            await self._write_profile_attributes(
                session, tenant_id=tenant_id, product_id=int(product_id),
                job_uuid=job_uuid, row_number=row["source_row_number"], values=values,
            )
            relation_items = self._relations_from_values(values)
            if relation_items:
                await self.products.replace_relations(
                    session, tenant_id=tenant_id, user_id=user_id,
                    product_id=int(product_id), items=relation_items,
                )
            await session.execute(text("""
                UPDATE product_import_rows SET imported_product_id=:product_id
                 WHERE tenant_id=:tenant_id AND job_id=:job_id
                   AND source_row_number=:source_row_number
            """), {
                "tenant_id": tenant_id, "job_id": job["id"],
                "source_row_number": row["source_row_number"],
                "product_id": product_id,
            })

        aliases = job["schema_snapshot"].get("aliases", [])
        for alias in aliases:
            target = await session.scalar(text("""
                SELECT sku FROM products
                 WHERE tenant_id=:tenant_id AND deleted_at IS NULL
                   AND sku_compare_key=:target_key
            """), {"tenant_id": tenant_id,
                     "target_key": _compare_key(alias["product_sku"])})
            if target is None:
                raise BusinessError("PRODUCT_IMPORT_ALIAS_STALE", "别名目标产品已变化", status_code=409)
            await session.execute(text("""
                INSERT INTO tenant_forecast_sku_aliases(
                  tenant_id,source_context,product_sku,source_sku
                ) VALUES(:tenant_id,:context,:product_sku,:source_sku)
                ON CONFLICT(tenant_id,source_context,product_sku) DO UPDATE
                  SET source_sku=EXCLUDED.source_sku
            """), {
                "tenant_id": tenant_id, "context": contract_context(tenant_id),
                "product_sku": target, "source_sku": alias["source_sku"],
            })

        committed_at = datetime.now().astimezone()
        next_snapshot = job["schema_snapshot"] | {
            "created_rows": created_rows,
            "updated_rows": updated_rows,
            "skipped_rows": skipped_rows,
            "alias_rows": len(aliases),
        }
        imported_rows = created_rows + updated_rows
        await session.execute(text("""
            UPDATE product_import_jobs
               SET status='completed',imported_rows=:imported_rows,
                   committed_at=:committed_at,schema_snapshot=CAST(:snapshot AS jsonb)
             WHERE id=:job_id AND tenant_id=:tenant_id
        """), {
            "tenant_id": tenant_id, "job_id": job["id"],
            "imported_rows": imported_rows, "committed_at": committed_at,
            "snapshot": _stable_json(next_snapshot),
        })
        await self.audit.record(
            session, tenant_id=tenant_id, actor_user_id=user_id,
            action_code="product.import.commit", resource_type="product_import_job",
            resource_id=job["id"], request_id=request_id,
            before={"status": "ready", "preview_sha256": preview_sha256},
            after={
                "status": "completed", "source_sha256": job["source_sha256"],
                "imported_rows": imported_rows, "created_rows": created_rows,
                "updated_rows": updated_rows, "skipped_rows": skipped_rows,
                "alias_rows": len(aliases),
            },
        )
        return {
            "job_uuid": job_uuid, "status": "completed",
            "imported_rows": imported_rows, "created_rows": created_rows,
            "updated_rows": updated_rows, "skipped_rows": skipped_rows,
            "alias_rows": len(aliases),
            "committed_at": committed_at,
        }

    async def _write_profile_attributes(
        self, session: AsyncSession, *, tenant_id: int, product_id: int,
        job_uuid: str, row_number: int, values: dict[str, Any],
    ) -> None:
        definitions = [
            ("overall_length", values.get("length"), values.get("dimension_unit")),
            ("overall_width", values.get("width"), values.get("dimension_unit")),
            ("overall_height", values.get("height"), values.get("dimension_unit")),
            ("weight", values.get("weight"), values.get("weight_unit")),
            ("moq", values.get("moq"), "件"),
            ("factory_price", values.get("factory_price"), values.get("currency")),
        ]
        attributes = [item for item in definitions if item[1] is not None]
        if not attributes:
            return
        profile_id = await self.products.ensure_draft_profile(
            session, tenant_id=tenant_id, product_id=product_id
        )
        for code, value, unit in attributes:
            await session.execute(text("""
                INSERT INTO product_attributes(
                  tenant_id,profile_version_id,attribute_code,value,unit,source_type,
                  source_locator,confidence,confirmation_status
                ) VALUES(
                  :tenant_id,:profile_id,:code,CAST(:value AS jsonb),:unit,'user_input',
                  CAST(:locator AS jsonb),1,'unconfirmed'
                )
                ON CONFLICT(profile_version_id,attribute_code) DO UPDATE SET
                  value=EXCLUDED.value,unit=EXCLUDED.unit,source_type='user_input',
                  source_locator=EXCLUDED.source_locator,confidence=1,
                  confirmation_status='unconfirmed'
            """), {
                "tenant_id": tenant_id, "profile_id": profile_id, "code": code,
                "value": json.dumps(float(Decimal(str(value)))),
                "unit": unit,
                "locator": _stable_json({
                    "input": "product_import", "job_uuid": job_uuid,
                    "source_row_number": row_number,
                }),
            })
        await session.execute(text("""
            UPDATE products SET current_profile_version_id=:profile_id,
                   analysis_status='profile_pending'
             WHERE id=:product_id AND tenant_id=:tenant_id
        """), {
            "profile_id": profile_id, "product_id": product_id, "tenant_id": tenant_id,
        })

    @staticmethod
    def _relations_from_values(values: dict[str, Any]) -> list[dict[str, Any]]:
        output = []
        if values.get("spu_code"):
            output.append({
                "group_type": "spu", "group_code": values["spu_code"],
                "group_name": values["spu_code"], "member_role": "variant",
                "quantity": Decimal("1"),
            })
        if values.get("bundle_code"):
            output.append({
                "group_type": "bundle", "group_code": values["bundle_code"],
                "group_name": values["bundle_code"], "member_role": "item",
                "quantity": Decimal(values.get("bundle_quantity", "1")),
            })
        if values.get("bom_code"):
            output.append({
                "group_type": "bom", "group_code": values["bom_code"],
                "group_name": values["bom_code"], "member_role": "component",
                "quantity": Decimal(values.get("bom_quantity", "1")),
            })
        return output

    async def cancel(
        self, session: AsyncSession, *, tenant_id: int, user_id: int,
        job_uuid: str, request_id: str,
    ) -> dict[str, Any]:
        job = await self._job_row(
            session, tenant_id=tenant_id, job_uuid=job_uuid, lock=True
        )
        if job["status"] in {"completed", "cancelled"}:
            if job["status"] == "completed":
                raise BusinessError("PRODUCT_IMPORT_IMMUTABLE", "已完成任务不能取消", status_code=409)
            return await self.job(
                session, tenant_id=tenant_id, job_uuid=job_uuid, include_rows=False
            )
        await session.execute(text("""
            UPDATE product_import_jobs SET status='cancelled',cancelled_at=CURRENT_TIMESTAMP
             WHERE id=:job_id AND tenant_id=:tenant_id
        """), {"job_id": job["id"], "tenant_id": tenant_id})
        await self.audit.record(
            session, tenant_id=tenant_id, actor_user_id=user_id,
            action_code="product.import.cancel", resource_type="product_import_job",
            resource_id=job["id"], request_id=request_id,
            before={"status": job["status"]}, after={"status": "cancelled"},
        )
        return await self.job(
            session, tenant_id=tenant_id, job_uuid=job_uuid, include_rows=False
        )

    async def error_report(
        self, session: AsyncSession, *, tenant_id: int, job_uuid: str
    ) -> bytes:
        job = await self._job_row(
            session, tenant_id=tenant_id, job_uuid=job_uuid
        )
        rows = (await session.execute(text("""
            SELECT source_row_number,source_values,normalized_values,
                   validation_errors,validation_warnings
              FROM product_import_rows
             WHERE tenant_id=:tenant_id AND job_id=:job_id
               AND (
                 jsonb_array_length(validation_errors)>0
                 OR jsonb_array_length(validation_warnings)>0
               )
             ORDER BY source_row_number
        """), {"tenant_id": tenant_id, "job_id": job["id"]})).mappings().all()
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "错误明细"
        source_headers = sorted({
            key for row in rows for key in row["source_values"]
        })
        sheet.append([
            "原行号", *source_headers, "错误码", "错误字段", "错误详情",
            "修复建议", "警告详情", "标准化预览",
        ])
        for row in rows:
            errors = row["validation_errors"]
            warnings = row["validation_warnings"]
            sheet.append([
                row["source_row_number"],
                *[row["source_values"].get(header) for header in source_headers],
                "\n".join(item["code"] for item in errors),
                "\n".join(item["field"] for item in errors),
                "\n".join(item["message"] for item in errors),
                "\n".join(item["suggestion"] for item in errors),
                "\n".join(item["message"] for item in warnings),
                _stable_json(row["normalized_values"]),
            ])
        summary = workbook.create_sheet("错误统计")
        summary.append(["job_uuid", job_uuid])
        summary.append(["source_sha256", job["source_sha256"]])
        summary.append(["total_rows", job["total_rows"]])
        summary.append(["error_rows", job["error_rows"]])
        summary.append(["warning_rows", job["warning_rows"]])
        counts: dict[str, int] = {}
        for row in rows:
            for issue in row["validation_errors"]:
                counts[issue["code"]] = counts.get(issue["code"], 0) + 1
        summary.append([])
        summary.append(["错误码", "数量"])
        for code, count in sorted(counts.items()):
            summary.append([code, count])
        for target in workbook.worksheets:
            target.freeze_panes = "A2"
            for cell in target[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="B64C3C")
            for column in target.columns:
                target.column_dimensions[column[0].column_letter].width = min(
                    42, max(12, max(len(str(cell.value or "")) for cell in column) + 2)
                )
        buffer = BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()
