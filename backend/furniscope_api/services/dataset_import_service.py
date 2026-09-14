"""JSON/CSV/Excel market import worker for authorized enterprise packs."""

from datetime import datetime, timezone
import csv
import hashlib
from io import BytesIO
from io import StringIO
import json
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncSession

from ..config import ApiSettings
from ..errors import BusinessError
from .demo_storage import DemoStorage
from .idempotency_service import IdempotencyService
from .audit_service import AuditService
from .review_cleaning import (
    DUPLICATE,
    EMPTY,
    GARBLED,
    REASON_LABELS,
    SPAM,
    content_hash,
    infer_sentiment,
    review_invalid_reason,
    unique_invalid_hash,
)


class DatasetImportService:
    demo_only = False

    LISTING_FIELDS = {
        "商品编号": "platform_listing_id", "platform_listing_id": "platform_listing_id",
        "商品标题": "title", "title": "title",
        "币种": "currency", "currency": "currency",
        "售价": "sale_price", "sale_price": "sale_price",
        "划线价": "list_price", "list_price": "list_price", "原价": "list_price",
        "采集时间": "captured_at", "captured_at": "captured_at",
        "上架日期": "first_available_date", "first_available_date": "first_available_date",
        "品牌": "brand", "brand": "brand",
        "评分": "rating", "rating": "rating",
        "评论数": "review_count", "review_count": "review_count",
        "品类": "category_code", "category_code": "category_code",
        "目标市场": "market_country", "market_country": "market_country", "市场": "market_country",
    }
    REVIEW_FIELDS = {
        "商品编号": "platform_listing_id", "platform_listing_id": "platform_listing_id",
        "评价编号": "platform_review_id", "platform_review_id": "platform_review_id",
        "评分": "rating", "rating": "rating",
        "评价内容": "content_original", "content_original": "content_original",
        "语言": "language_code", "language_code": "language_code",
        "是否已购买": "verified_purchase", "verified_purchase": "verified_purchase",
        "评价时间": "reviewed_at", "reviewed_at": "reviewed_at",
        "用户属地": "reviewer_location", "reviewer_location": "reviewer_location",
        "属地": "reviewer_location", "location": "reviewer_location",
        "情感": "sentiment", "sentiment": "sentiment",
    }

    def __init__(self, settings: ApiSettings) -> None:
        self.storage=DemoStorage(settings)
        self.idempotency=IdempotencyService()
        self.audit=AuditService()

    async def accept(self,session:AsyncSession,*,tenant_id:int,user_id:int,dataset_id:int,
                     idempotency_key:str,deduplication_strategy:str,filename:str,mime:str,content:bytes,
                     response_envelope,request_id:str) -> tuple[int,dict[str,Any]]:
        if deduplication_strategy != "platform_id_latest":
            raise BusinessError("FIELD_MAPPING_INVALID","P0去重策略仅支持platform_id_latest",status_code=422)
        dataset=(await session.execute(text("""
            SELECT id,status FROM market_datasets WHERE id=:id AND tenant_id=:tenant AND deleted_at IS NULL FOR UPDATE
        """),{"id":dataset_id,"tenant":tenant_id})).mappings().one_or_none()
        if dataset is None: raise BusinessError("DATASET_NOT_FOUND","数据集不存在或不可访问",status_code=404)
        if dataset["status"] == "validating": raise BusinessError("DATASET_IMPORT_IN_PROGRESS","数据集正在导入",status_code=409)
        try:
            parsed=self.parse_payload(filename=filename,mime=mime,content=content)
            if (not isinstance(parsed,dict) or not isinstance(parsed.get("listings"),list)
                    or not isinstance(parsed.get("reviews"),list)
                    or not parsed["listings"] or not parsed["reviews"]):
                raise ValueError
        except (UnicodeDecodeError,json.JSONDecodeError,ValueError,KeyError) as exc:
            raise BusinessError("DATASET_FILE_INVALID","数据文件无法识别，请下载模板并检查必填内容",status_code=422) from exc
        decision=await self.idempotency.begin(session,tenant_id=tenant_id,actor_user_id=user_id,
            route_code="API-DAT-02",http_method="POST",idempotency_key=idempotency_key,
            request_payload={"dataset_id":dataset_id,"deduplication_strategy":deduplication_strategy,
                             "file":{"name":filename,"mime":mime,"sha256":hashlib.sha256(content).hexdigest()}})
        if decision.action=="replay": return int(decision.response_status),dict(decision.response_body or {})
        key,digest=self.storage.store(tenant_id=tenant_id,filename=filename,mime_type=mime,content=content)
        asset_type = {".xlsx": "spreadsheet", ".csv": "csv", ".json": "json"}.get(
            Path(filename).suffix.lower(), "other",
        )
        asset_id=int((await session.execute(text("""
            INSERT INTO file_assets(tenant_id,owner_type,owner_id,asset_type,original_filename,storage_key,mime_type,
              size_bytes,sha256,security_status,parse_status,created_by)
            VALUES(:tenant,'dataset',:dataset,:asset_type,:name,:key,:mime,:size,:sha,'clean','processing',:user) RETURNING id
        """),{"tenant":tenant_id,"dataset":dataset_id,"name":filename,"key":key,"mime":mime,
               "size":len(content),"sha":digest,"user":user_id,"asset_type":asset_type})).scalar_one())
        accepted_at=datetime.now(timezone.utc)
        await session.execute(text("""
            UPDATE market_datasets SET import_asset_id=:asset,status='validating' WHERE id=:id AND tenant_id=:tenant
        """),{"asset":asset_id,"id":dataset_id,"tenant":tenant_id})
        await self.audit.record(
            session, tenant_id=tenant_id, actor_user_id=user_id,
            action_code="dataset.import.accept", resource_type="market_dataset",
            resource_id=dataset_id, request_id=request_id,
            after={"asset_id": asset_id, "filename": filename, "mime": mime,
                   "size_bytes": len(content), "sha256": digest,
                   "deduplication_strategy": deduplication_strategy, "status": "validating"},
        )
        envelope=response_envelope({"dataset_id":dataset_id,"status":"validating","accepted_at":accepted_at})
        await self.idempotency.finish(session,tenant_id=tenant_id,record_id=decision.record_id,response_status=202,
            response_body=envelope,resource_type="market_datasets",resource_public_id=str(dataset_id))
        return 202,envelope

    @classmethod
    def parse_payload(cls, *, filename: str, mime: str, content: bytes) -> dict[str, list[dict[str, Any]]]:
        suffix = Path(filename).suffix.lower()
        if suffix == ".json" or mime == "application/json":
            payload = json.loads(content)
            if not isinstance(payload, dict):
                raise ValueError("数据文件内容无效")
            return payload
        if suffix == ".csv" or mime in {"text/csv", "application/csv"}:
            rows = list(csv.DictReader(StringIO(content.decode("utf-8-sig"))))
            listings: list[dict[str, Any]] = []
            reviews: list[dict[str, Any]] = []
            for source in rows:
                row = {str(key).strip(): value for key, value in source.items() if key and value not in (None, "")}
                record_type = str(row.get("record_type") or row.get("数据类型") or row.get("type") or "").lower()
                is_review = record_type in {"review", "评论", "评价"} or any(key in row for key in ("评价编号", "platform_review_id", "评价内容", "content_original"))
                mapping = cls.REVIEW_FIELDS if is_review else cls.LISTING_FIELDS
                normalized = {mapping[key]: value for key, value in row.items() if key in mapping}
                (reviews if is_review else listings).append(normalized)
            if not listings or not reviews:
                raise ValueError("CSV 必须同时包含 record_type=listing 与 record_type=review 的数据行")
            return {"listings": listings, "reviews": reviews}
        if suffix != ".xlsx":
            raise ValueError("仅支持市场数据工作簿")
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        try:
            listings: list[dict[str, Any]] = []
            reviews: list[dict[str, Any]] = []
            for sheet in workbook.worksheets:
                rows = sheet.iter_rows(values_only=True)
                headers = next(rows, None)
                if not headers:
                    continue
                listing_map = {index: cls.LISTING_FIELDS.get(str(value).strip()) for index, value in enumerate(headers) if value is not None}
                review_map = {index: cls.REVIEW_FIELDS.get(str(value).strip()) for index, value in enumerate(headers) if value is not None}
                is_listing = {"platform_listing_id", "title", "currency", "sale_price", "captured_at"}.issubset(set(listing_map.values()))
                is_review = {"platform_listing_id", "platform_review_id", "content_original"}.issubset(set(review_map.values()))
                target, mapping = (listings, listing_map) if is_listing else (reviews, review_map) if is_review else (None, None)
                if target is None:
                    continue
                for values in rows:
                    row = {field: values[index] for index, field in mapping.items() if field and index < len(values) and values[index] not in (None, "")}
                    if not row:
                        continue
                    for date_field in ("captured_at", "reviewed_at", "first_available_date"):
                        if isinstance(row.get(date_field), datetime):
                            row[date_field] = row[date_field].isoformat()
                    if "verified_purchase" in row:
                        row["verified_purchase"] = str(row["verified_purchase"]).strip().lower() in {"1", "true", "yes", "是", "已购买"}
                    target.append(row)
            if not listings or not reviews:
                raise ValueError("工作簿必须包含商品信息和消费者评价")
            return {"listings": listings, "reviews": reviews}
        finally:
            workbook.close()

    async def run_import(self,session:AsyncSession,*,tenant_id:int,dataset_id:int,content:bytes,
                         filename: str | None = None, mime: str | None = None) -> None:
        if not filename or not mime:
            asset = (await session.execute(text("""
                SELECT original_filename,mime_type FROM file_assets
                 WHERE tenant_id=:tenant AND owner_type='dataset' AND owner_id=:dataset
                 ORDER BY id DESC LIMIT 1
            """), {"tenant": tenant_id, "dataset": dataset_id})).mappings().one()
            filename, mime = asset["original_filename"], asset["mime_type"]
        payload=self.parse_payload(filename=filename,mime=mime,content=content)
        dataset= (await session.execute(text("""
            SELECT category_code,market_country FROM market_datasets WHERE id=:dataset AND tenant_id=:tenant
        """),{"dataset":dataset_id,"tenant":tenant_id})).mappings().one_or_none()
        if not dataset:
            raise ValueError("dataset category unavailable")
        await self._ensure_preview_columns(session)
        await self._replace_dataset_rows(session, tenant_id=tenant_id, dataset_id=dataset_id)
        listing_ids = await self._insert_listings(
            session, tenant_id=tenant_id, dataset_id=dataset_id,
            dataset_category=dataset["category_code"], dataset_country=dataset["market_country"],
            listings=payload["listings"],
        )
        review_stats = await self._insert_reviews(
            session, tenant_id=tenant_id, dataset_id=dataset_id,
            listing_ids=listing_ids, reviews=payload["reviews"],
        )
        raw_listings = len(payload["listings"])
        raw_reviews = len(payload["reviews"])
        kept_listings = len(listing_ids)
        valid_reviews = review_stats["valid"]
        quality = 0
        if kept_listings:
            retention = (valid_reviews / raw_reviews) if raw_reviews else 1
            quality = min(100, round(20 + retention * 80))
        report = {
            "source": "uploaded_market_file",
            "schema_valid": True,
            "overwrite": True,
            "raw_listing_count": raw_listings,
            "kept_listing_count": kept_listings,
            "raw_review_count": raw_reviews,
            "valid_review_count": valid_reviews,
            "filtered_review_count": max(0, raw_reviews - valid_reviews),
            "filter_reasons": review_stats["reasons"],
        }
        await session.execute(text("""
            UPDATE market_datasets SET status='ready',listing_count=:listings,review_count=:reviews,
              valid_review_count=:valid,quality_score=:quality,
              quality_report=CAST(:report AS jsonb),limitations='[]'::jsonb,
              version_no=version_no+1,updated_at=CURRENT_TIMESTAMP
             WHERE id=:dataset AND tenant_id=:tenant
        """),{"listings":kept_listings,"reviews":raw_reviews,"valid":valid_reviews,"quality":quality,
              "report":json.dumps(report,ensure_ascii=False),"dataset":dataset_id,"tenant":tenant_id})
        await session.execute(text("UPDATE file_assets SET parse_status='succeeded' WHERE tenant_id=:tenant AND owner_type='dataset' AND owner_id=:dataset"),{"tenant":tenant_id,"dataset":dataset_id})
        from .competitor_tracking import CompetitorTrackingService
        await CompetitorTrackingService().capture_dataset(
            session, tenant_id=tenant_id, dataset_id=dataset_id, source="dataset_import",
        )

    @staticmethod
    async def _ensure_preview_columns(session: AsyncSession) -> None:
        await session.execute(text("ALTER TABLE reviews ADD COLUMN IF NOT EXISTS reviewer_location VARCHAR(100)"))
        await session.execute(text("ALTER TABLE reviews ADD COLUMN IF NOT EXISTS sentiment VARCHAR(16)"))

    @staticmethod
    async def _replace_dataset_rows(session: AsyncSession, *, tenant_id: int, dataset_id: int) -> None:
        params = {"tenant": tenant_id, "dataset": dataset_id}
        await session.execute(text("""
            DELETE FROM evidence_links WHERE tenant_id=:tenant AND (
              (evidence_type='listing' AND evidence_id IN (
                  SELECT id FROM market_listings WHERE tenant_id=:tenant AND dataset_id=:dataset))
              OR (evidence_type='review_aspect' AND evidence_id IN (
                  SELECT a.id FROM review_aspects a JOIN reviews r ON r.id=a.review_id
                   WHERE r.tenant_id=:tenant AND r.dataset_id=:dataset))
            )
        """), params)
        await session.execute(text("""
            DELETE FROM review_aspects WHERE tenant_id=:tenant AND review_id IN (
              SELECT id FROM reviews WHERE tenant_id=:tenant AND dataset_id=:dataset)
        """), params)
        await session.execute(text("""
            DELETE FROM competitor_matches WHERE tenant_id=:tenant AND listing_id IN (
              SELECT id FROM market_listings WHERE tenant_id=:tenant AND dataset_id=:dataset)
        """), params)
        await session.execute(text("DELETE FROM reviews WHERE tenant_id=:tenant AND dataset_id=:dataset"), params)
        await session.execute(text("DELETE FROM market_listings WHERE tenant_id=:tenant AND dataset_id=:dataset"), params)

    @staticmethod
    def _parse_datetime(value: Any) -> datetime | None:
        if value in (None, ""):
            return None
        if isinstance(value, datetime):
            return value
        try:
            return datetime.fromisoformat(str(value).replace("Z", "+00:00"))
        except (TypeError, ValueError):
            return None

    async def _insert_listings(self, session: AsyncSession, *, tenant_id: int, dataset_id: int,
                               dataset_category: str, dataset_country: str,
                               listings: list[dict[str, Any]]) -> dict[str, int]:
        listing_ids: dict[str, int] = {}
        latest: dict[str, dict[str, Any]] = {}
        for item in listings:
            required = {"platform_listing_id", "title", "currency", "sale_price", "captured_at"}
            if not required.issubset(item):
                continue
            try:
                captured_at = datetime.fromisoformat(str(item["captured_at"]).replace("Z", "+00:00"))
            except (TypeError, ValueError):
                continue
            platform_id = str(item["platform_listing_id"])
            current = latest.get(platform_id)
            if current is None or captured_at >= current["captured_at"]:
                latest[platform_id] = {**item, "captured_at": captured_at}
        for platform_id, item in latest.items():
            attrs = dict(item.get("normalized_attributes") or {})
            market = str(item.get("market_country") or attrs.get("market_country") or dataset_country).upper()
            attrs["market_country"] = market
            first_available = str(item["first_available_date"])[:10] if item.get("first_available_date") else None
            list_price = item.get("list_price")
            sale_price = item.get("sale_price")
            try:
                list_price = None if list_price in (None, "") else float(list_price)
                sale_price = float(sale_price)
            except (TypeError, ValueError):
                continue
            row = (await session.execute(text("""
                INSERT INTO market_listings(tenant_id,dataset_id,platform_listing_id,title,category_code,currency,
                  list_price,sale_price,brand,rating,review_count,captured_at,first_available_date,
                  normalized_attributes,image_urls,raw_payload)
                VALUES(:tenant,:dataset,:platform_id,:title,:category,:currency,:list_price,:price,:brand,:rating,
                       :review_count,:captured,:first_available,CAST(:attrs AS jsonb),CAST(:images AS jsonb),CAST(:raw AS jsonb))
                RETURNING id
            """),{"tenant":tenant_id,"dataset":dataset_id,"platform_id":platform_id,
                   "title":str(item["title"]),"category":str(item.get("category_code") or dataset_category),
                   "currency":str(item["currency"]),"list_price":list_price,
                   "price":sale_price,"brand":item.get("brand"),"rating":item.get("rating"),
                   "review_count":item.get("review_count"),"captured":item["captured_at"],
                   "first_available":first_available,
                   "attrs":json.dumps(attrs,ensure_ascii=False),
                   "images":json.dumps(item.get("image_urls",[])),
                   "raw":json.dumps({key: (value.isoformat() if isinstance(value, datetime) else value)
                                     for key, value in item.items()}, ensure_ascii=False)})).scalar_one()
            listing_ids[platform_id] = int(row)
        return listing_ids

    async def _insert_reviews(self, session: AsyncSession, *, tenant_id: int, dataset_id: int,
                              listing_ids: dict[str, int], reviews: list[dict[str, Any]]) -> dict[str, Any]:
        seen: set[str] = set()
        reasons = {EMPTY: 0, DUPLICATE: 0, GARBLED: 0, SPAM: 0, "orphan": 0}
        valid = 0
        stored_ids: set[str] = set()
        for item in reviews:
            listing_id = listing_ids.get(str(item.get("platform_listing_id") or ""))
            review_id = str(item.get("platform_review_id") or "").strip()
            content_original = str(item.get("content_original") or "")
            if not listing_id or not review_id or review_id in stored_ids:
                reasons["orphan" if not listing_id or not review_id else DUPLICATE] += 1
                continue
            reason = review_invalid_reason(content_original, seen)
            is_valid = reason is None
            if reason:
                reasons[reason] += 1
            else:
                valid += 1
            digest = content_hash(content_original) if is_valid else unique_invalid_hash(review_id, content_original)
            sentiment = str(item.get("sentiment") or "").lower()
            if sentiment not in {"positive", "neutral", "negative"}:
                sentiment = infer_sentiment(content_original, item.get("rating"))
            rating = item.get("rating")
            try:
                rating = None if rating in (None, "") else float(rating)
            except (TypeError, ValueError):
                rating = None
            reviewed_at = self._parse_datetime(item.get("reviewed_at"))
            await session.execute(text("""
                INSERT INTO reviews(tenant_id,dataset_id,listing_id,platform_review_id,rating,content_original,
                  language_code,reviewed_at,verified_purchase,content_hash,is_valid,invalid_reason,
                  reviewer_location,sentiment)
                VALUES(:tenant,:dataset,:listing,:review_id,:rating,:content,:language,:reviewed_at,:verified,:hash,
                       :is_valid,:invalid_reason,:location,:sentiment)
            """),{"tenant":tenant_id,"dataset":dataset_id,"listing":listing_id,"review_id":review_id,
                   "rating":rating,"content":content_original,
                   "language":item.get("language_code") or "en",
                   "reviewed_at":reviewed_at,
                   "verified":bool(item.get("verified_purchase", False)),
                   "hash":digest,"is_valid":is_valid,
                   "invalid_reason": None if is_valid else REASON_LABELS.get(reason, reason),
                   "location": item.get("reviewer_location") or item.get("location") or None,
                   "sentiment": sentiment})
            stored_ids.add(review_id)
        return {"valid": valid, "reasons": {key: value for key, value in reasons.items() if value}}

    async def run_demo(self,session:AsyncSession,*,tenant_id:int,dataset_id:int,content:bytes) -> None:
        await self.run_import(session, tenant_id=tenant_id, dataset_id=dataset_id,
                              content=content, filename="legacy-market-data.json",
                              mime="application/json")
