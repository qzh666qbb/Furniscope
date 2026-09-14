from __future__ import annotations

import json
from io import BytesIO

import httpx
from fastapi.testclient import TestClient
from openpyxl import Workbook
import pytest

from furniscope_api.app import create_app
from furniscope_api.config import ApiSettings
from furniscope_api.services.document_parser import DocumentParseError, DocumentParser
from furniscope_api.services.dataset_import_service import DatasetImportService
from furniscope_api.services.review_cleaning import infer_sentiment, review_invalid_reason
from furniscope_api.services.model_router_client import ServiceModelRouterClient


class NoDatabase:
    async def close(self):
        return None


def test_xlsx_document_parser_extracts_real_cells_and_rejects_spoofed_files() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "产品参数"
    sheet.append(["SKU", "FS-AUTH-001"])
    sheet.append(["材质", "橡木"])
    payload = BytesIO()
    workbook.save(payload)
    parser = DocumentParser(max_bytes=1_000_000)
    parsed = parser.parse("product.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", payload.getvalue())
    assert "FS-AUTH-001" in parsed.text and "橡木" in parsed.text
    assert parsed.extraction_method == "xlsx_cells"
    with pytest.raises(DocumentParseError, match="XLSX") as error:
        parser.parse("spoof.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", b"not-a-zip")
    assert error.value.code == "FILE_SIGNATURE_INVALID"


@pytest.mark.asyncio
async def test_bailian_embedding_and_rerank_contracts() -> None:
    calls: list[httpx.Request] = []

    def handler(request: httpx.Request) -> httpx.Response:
        calls.append(request)
        if request.url.path.endswith("/embeddings"):
            payload = json.loads(request.content)
            return httpx.Response(200, json={
                "model": payload["model"],
                "data": [
                    {"index": 1, "embedding": [0.2] * 1024},
                    {"index": 0, "embedding": [0.1] * 1024},
                ],
                "usage": {"total_tokens": 12},
            })
        payload = json.loads(request.content)
        return httpx.Response(200, json={
            "model": payload["model"],
            "results": [
                {"index": 1, "relevance_score": 0.91},
                {"index": 0, "relevance_score": 0.12},
            ],
            "usage": {"total_tokens": 8},
        })

    settings = ApiSettings(
        database_url="postgresql+asyncpg://postgres@127.0.0.1:5432/furniscope",
        aliyun_model_router_api_key="test-key",
        internal_service_token="test-token",
    )
    async with httpx.AsyncClient(transport=httpx.MockTransport(handler)) as http_client:
        client = ServiceModelRouterClient(settings, client=http_client)
        embedding = await client.embeddings(["oak dining table", "fabric sofa"])
        ranking = await client.rerank("oak table", ["fabric sofa", "oak table"])

    assert embedding["provider"] == "aliyun_bailian"
    assert embedding["dimensions"] == 1024
    assert embedding["vectors"][0][0] == 0.1
    assert ranking["provider"] == "aliyun_bailian"
    assert ranking["ranking"][0] == {"index": 1, "score": 0.91, "rank": 1}
    assert [request.url.path.rsplit("/", 1)[-1] for request in calls] == ["embeddings", "reranks"]


@pytest.mark.asyncio
async def test_bailian_embedding_batches_more_than_ten_texts() -> None:
    sizes: list[int] = []

    def handler(request: httpx.Request) -> httpx.Response:
        payload = json.loads(request.content)
        texts = payload["input"]
        sizes.append(len(texts))
        return httpx.Response(200, json={
            "model": payload["model"],
            "data": [{"index": index, "embedding": [0.1] * 1024} for index in range(len(texts))],
            "usage": {"total_tokens": len(texts)},
        })

    settings = ApiSettings(
        database_url="postgresql+asyncpg://postgres@127.0.0.1:5432/furniscope",
        aliyun_model_router_api_key="test-key",
        internal_service_token="test-token",
    )
    texts = [f"furniture listing {index}" for index in range(12)]
    async with httpx.AsyncClient(transport=httpx.MockTransport(handler)) as http_client:
        client = ServiceModelRouterClient(settings, client=http_client)
        embedding = await client.embeddings(texts)
        with pytest.raises(ValueError, match="1-10 non-empty texts"):
            await client.embeddings(["oak", "  "])

    assert sizes == [10, 2]
    assert len(embedding["vectors"]) == 12


def test_market_workbook_parses_business_facing_sheets() -> None:
    workbook = Workbook()
    listings = workbook.active
    listings.title = "商品信息"
    listings.append(["商品编号", "商品标题", "品牌", "售价", "币种", "评分", "评论数", "采集时间"])
    listings.append(["HF-MKT-01", "Modern accent chair", "HeFeng", 189, "USD", 4.5, 20, "2026-09-01"])
    reviews = workbook.create_sheet("消费者评价")
    reviews.append(["商品编号", "评价编号", "评分", "评价内容", "语言", "是否已购买", "评价时间"])
    reviews.append(["HF-MKT-01", "REV-01", 5, "The chair is comfortable.", "en", "是", "2026-09-02"])
    payload = BytesIO()
    workbook.save(payload)

    parsed = DatasetImportService.parse_payload(
        filename="market-data.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        content=payload.getvalue(),
    )

    assert parsed["listings"][0]["platform_listing_id"] == "HF-MKT-01"
    assert parsed["listings"][0]["sale_price"] == 189
    assert parsed["reviews"][0]["content_original"] == "The chair is comfortable."
    assert parsed["reviews"][0]["verified_purchase"] is True


def test_market_csv_parses_price_range_and_reviewer_location() -> None:
    csv_text = (
        "record_type,platform_listing_id,title,currency,sale_price,list_price,captured_at,"
        "market_country,platform_review_id,rating,content_original,reviewer_location\n"
        "listing,HF-MKT-02,Compact sofa,USD,199,249,2026-09-01T00:00:00Z,US,,,,\n"
        "review,HF-MKT-02,,,,,,,REV-02,2,The frame is wobbly and I am disappointed.,California\n"
    )
    parsed = DatasetImportService.parse_payload(
        filename="market-data.csv", mime="text/csv", content=csv_text.encode()
    )
    assert parsed["listings"][0]["list_price"] == "249"
    assert parsed["listings"][0]["market_country"] == "US"
    assert parsed["reviews"][0]["reviewer_location"] == "California"


def test_review_cleaning_filters_empty_garbled_spam_and_duplicates() -> None:
    seen: set[str] = set()
    assert review_invalid_reason("   ", seen) == "empty"
    assert review_invalid_reason("...", seen) == "empty"
    assert review_invalid_reason("\ufffd\ufffd\ufffd%%%@@@", seen) == "garbled"
    assert review_invalid_reason("good nice", seen) == "spam"
    assert review_invalid_reason("The seat is comfortable and sturdy for daily use.", seen) is None
    assert review_invalid_reason("The seat is comfortable and sturdy for daily use.", seen) == "duplicate"


def test_review_sentiment_prefers_text_over_star_rating() -> None:
    assert infer_sentiment("The chair is comfortable and I love it.", 2) == "positive"
    assert infer_sentiment("Very disappointed, the frame is broken.", 5) == "negative"
    assert infer_sentiment("It arrived on time.", 5) == "positive"


def test_competitor_snapshot_diff_detects_price_title_promo_and_bullets() -> None:
    from furniscope_api.services.competitor_tracking import diff_snapshots, listing_fingerprint, promo_label
    previous = {
        "title": "Old title", "sale_price": 199, "list_price": 249,
        "image_urls": ["a.jpg"], "bullet_points": ["one"], "promo_label": "直降 20%",
    }
    current = {
        "title": "New title", "sale_price": 179, "list_price": 249,
        "image_urls": ["a.jpg"], "bullet_points": ["one", "two"], "promo_label": "直降 28%",
    }
    types = {item["change_type"] for item in diff_snapshots(previous, current)}
    assert types == {"price", "title", "bullets", "promo"}
    assert listing_fingerprint(previous) != listing_fingerprint(current)
    assert promo_label(179, 249) == "直降 28%"


def test_internal_model_endpoints_require_service_token(monkeypatch) -> None:
    settings = ApiSettings(
        app_env="test",
        database_url="postgresql+asyncpg://test:test@localhost/test",
        internal_service_token="test-internal-token",
        aliyun_model_router_api_key="test-model-key",
    )
    app = create_app(settings, NoDatabase())

    async def structured(_self, **_kwargs):
        return {"answer": "validated"}

    monkeypatch.setattr(ServiceModelRouterClient, "structured", structured)

    async def embeddings(_self, texts, dimensions=None):
        size = dimensions or 1024
        return {"vectors": [[0.1] * size], "dimensions": size, "provider": "aliyun_bailian"}

    async def rerank(_self, _query, _candidates):
        return {"ranking": [{"index": 1, "score": 0.91, "rank": 1}]}

    monkeypatch.setattr(ServiceModelRouterClient, "embeddings", embeddings)
    monkeypatch.setattr(ServiceModelRouterClient, "rerank", rerank)
    with TestClient(app) as client:
        denied = client.post("/internal/v1/model-router/embeddings", json={"texts": ["oak"]})
        headers = {"X-Internal-Token": "test-internal-token"}
        rejected_dim = client.post("/internal/v1/model-router/embeddings", headers=headers,
                                   json={"texts": ["oak"], "dimensions": 32})
        embedded = client.post("/internal/v1/model-router/embeddings", headers=headers,
                               json={"texts": ["oak"], "dimensions": 1024})
        reranked = client.post("/internal/v1/model-router/rerank", headers=headers,
                               json={"query": "oak", "candidates": ["sofa", "oak"]})
        invoked = client.post("/internal/v1/model-router/invoke", headers=headers,
                              json={"messages": [{"role": "user", "content": "x"}],
                                    "schema": {"required": ["answer"]}})
    assert denied.status_code == 401
    assert rejected_dim.status_code == 422
    assert embedded.status_code == 200 and len(embedded.json()["data"]["vectors"][0]) == 1024
    assert reranked.status_code == 200 and reranked.json()["data"]["ranking"][0]["index"] == 1
    assert invoked.status_code == 200 and invoked.json()["data"]["output"] == {"answer": "validated"}
