"""Product master import contracts, performance, and PostgreSQL HTTP lifecycle."""

from __future__ import annotations

import asyncio
from hashlib import sha256
from io import BytesIO
import json
import os
from time import perf_counter
from uuid import uuid4

from openpyxl import Workbook, load_workbook
import pytest

from furniscope_api.services.product_import_service import (
    FIELD_SCHEMA,
    ProductImportService,
    TEMPLATE_VERSION,
    _normalize_values,
    build_template,
    suggest_mapping,
    template_schema_sha256,
)


TEST_DSN = os.getenv("FURNISCOPE_TEST_DATABASE_URL")
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _workbook_bytes(
    rows: list[list[object]], *, headers: list[str] | None = None,
    alias_rows: list[list[object]] | None = None,
) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Product_Master"
    sheet.append(headers or [field["title"] for field in FIELD_SCHEMA])
    for row in rows:
        sheet.append(row)
    aliases = workbook.create_sheet("SKU_Alias")
    aliases.append(["来源系统", "来源SKU", "产品SKU"])
    for row in alias_rows or []:
        aliases.append(row)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _issue_codes(issues: list[dict[str, str]]) -> set[str]:
    return {item["code"] for item in issues}


def test_dynamic_template_has_versioned_five_sheet_contract() -> None:
    content, file_sha, schema_sha = build_template()
    assert file_sha == sha256(content).hexdigest()
    assert schema_sha == template_schema_sha256()

    workbook = load_workbook(BytesIO(content), data_only=False)
    assert workbook.sheetnames == [
        "填写说明", "Product_Master", "SKU_Alias", "Dictionaries", "Examples",
    ]
    guide = {
        str(row[0].value): row[1].value
        for row in workbook["填写说明"].iter_rows(min_row=1, max_col=2)
        if row[0].value
    }
    assert guide["template_version"] == TEMPLATE_VERSION
    assert guide["schema_sha256"] == schema_sha
    assert workbook["Product_Master"].max_row == 1
    assert workbook["SKU_Alias"].max_row == 1
    assert workbook["Examples"].max_row >= 3
    assert len(workbook["Product_Master"].data_validations.dataValidation) == 5


def test_mapping_and_normalization_are_strict_and_deterministic() -> None:
    headers = ["产品编码", "产品名称", "所属品类", "长度", "长宽高单位"]
    assert suggest_mapping(headers) == {
        "sku": "产品编码",
        "name": "产品名称",
        "category_code": "所属品类",
        "length": "长度",
        "dimension_unit": "长宽高单位",
    }

    normalized, errors, warnings = _normalize_values(
        {
            "sku": "  Sofa-001 ",
            "name": " 模块沙发 ",
            "category_code": "沙发",
            "length": "2,250",
            "dimension_unit": "毫米",
        },
        dictionary_mapping={"category_code": {"沙发": "sofa"}},
        unit_mapping={"毫米": "mm"},
        sku_from_cell="  Sofa-001 ",
    )
    assert errors == []
    assert warnings == []
    assert normalized | {} == {
        "sku": "Sofa-001",
        "name": "模块沙发",
        "category_code": "sofa",
        "length": "2250",
        "dimension_unit": "mm",
        "lifecycle_status": "active",
    }

    _, missing, _ = _normalize_values({})
    assert [(item["code"], item["field"]) for item in missing] == [
        ("REQUIRED", "sku"),
        ("REQUIRED", "name"),
        ("REQUIRED", "category_code"),
    ]
    _, invalid, warning = _normalize_values({
        "sku": 10001,
        "name": "测试桌",
        "category_code": "desk",
        "length": -1,
        "weight": "not-a-number",
        "dimension_unit": "cm",
    }, sku_from_cell=10001)
    assert {"TEXT_TYPE", "ENUM_UNKNOWN", "NUMBER_INVALID"} <= _issue_codes(invalid)
    assert "UNIT_WITHOUT_VALUE" in _issue_codes(warning)

    _, unit_errors, category_warnings = _normalize_values({
        "sku": "TABLE-001",
        "name": "测试桌",
        "category_code": "table",
        "length": 120,
    })
    assert "UNIT_REQUIRED" in _issue_codes(unit_errors)
    assert "ANALYSIS_CATEGORY_UNVERIFIED" in _issue_codes(category_warnings)


def test_upload_inspection_detects_sheet_header_and_standardizes_known_values() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "产品数据"
    sheet.append(["导出说明", "该行不是表头"])
    sheet.append(["商品编码", "品名", "商品分类", "长", "尺寸单位", "币种"])
    sheet.append(["HF-001", "云感沙发", "沙发", 220, "厘米", "美元"])
    payload = BytesIO()
    workbook.save(payload)

    inspection = ProductImportService().inspect(
        filename="products.xlsx", mime_type=XLSX_MIME, content=payload.getvalue()
    )

    assert inspection["sheet_name"] == "产品数据"
    assert inspection["header_row"] == 2
    assert inspection["total_rows"] == 1
    assert inspection["field_mapping"] | {} == {
        "sku": "商品编码",
        "name": "品名",
        "category_code": "商品分类",
        "length": "长",
        "dimension_unit": "尺寸单位",
        "currency": "币种",
    }
    assert inspection["mapping_issues"] == []
    assert inspection["value_issues"] == []
    assert inspection["unit_mapping"] == {"厘米": "cm"}
    assert inspection["dictionary_mapping"] == {
        "category_code": {"沙发": "sofa"},
        "currency": {"美元": "USD"},
    }


def test_upload_inspection_only_returns_ambiguous_or_missing_required_fields() -> None:
    ambiguous = _workbook_bytes(
        [["HF-001", "HF-ALT-001", "扶手椅", "chair"]],
        headers=["SKU", "产品SKU", "产品名称", "品类代码"],
    )
    inspection = ProductImportService().inspect(
        filename="ambiguous.xlsx", mime_type=XLSX_MIME, content=ambiguous
    )
    assert "sku" not in inspection["field_mapping"]
    assert [(item["field_code"], item["kind"]) for item in inspection["mapping_issues"]] == [
        ("sku", "ambiguous")
    ]
    assert set(inspection["mapping_issues"][0]["candidates"]) == {"SKU", "产品SKU"}

    missing = _workbook_bytes(
        [["HF-001", "扶手椅"]],
        headers=["SKU", "产品名称"],
    )
    missing_inspection = ProductImportService().inspect(
        filename="missing.xlsx", mime_type=XLSX_MIME, content=missing
    )
    assert [(item["field_code"], item["kind"]) for item in missing_inspection["mapping_issues"]] == [
        ("category_code", "missing_required")
    ]


def test_case_insensitive_cross_row_duplicates_block_every_collision() -> None:
    rows = [
        {
            "source_row_number": number,
            "normalized_values": {
                "sku": sku, "name": f"产品 {number}",
                "category_code": "sofa", "lifecycle_status": "active",
            },
            "validation_errors": [],
            "validation_warnings": [],
        }
        for number, sku in ((2, " Sofa-01 "), (3, "sofa-01"))
    ]
    ProductImportService._apply_cross_row_rules(rows, {}, "create_only")
    assert all(row["planned_action"] == "invalid" for row in rows)
    assert all("SKU_DUPLICATE_FILE" in _issue_codes(row["validation_errors"]) for row in rows)


def test_create_only_skips_existing_sku_without_blocking_preflight() -> None:
    rows = [{
        "source_row_number": 2,
        "normalized_values": {
            "sku": " Sofa-01 ", "name": "已有沙发",
            "category_code": "sofa", "lifecycle_status": "active",
        },
        "validation_errors": [],
        "validation_warnings": [],
    }]

    ProductImportService._apply_cross_row_rules(
        rows, {"SOFA-01": 101}, "create_only"
    )

    assert rows[0]["planned_action"] == "skip"
    assert rows[0]["validation_errors"] == []
    assert "SKU_EXISTS_SKIPPED" in _issue_codes(rows[0]["validation_warnings"])


@pytest.mark.asyncio
async def test_error_report_preserves_source_row_error_code_and_repair_advice() -> None:
    service = ProductImportService()

    async def fake_job(*_args, **_kwargs):
        return {
            "id": 7,
            "source_sha256": "a" * 64,
            "total_rows": 1,
            "error_rows": 1,
            "warning_rows": 0,
        }

    class Result:
        def mappings(self):
            return self

        def all(self):
            return [{
                "source_row_number": 9,
                "source_values": {"SKU": "BAD-1", "品类": "desk"},
                "normalized_values": {"sku": "BAD-1"},
                "validation_errors": [{
                    "code": "ENUM_UNKNOWN",
                    "field": "category_code",
                    "message": "desk 不在受控词表中",
                    "suggestion": "使用 Dictionaries 中的标准代码",
                }],
                "validation_warnings": [],
            }]

    class Session:
        async def execute(self, *_args, **_kwargs):
            return Result()

    service._job_row = fake_job  # type: ignore[method-assign]
    content = await service.error_report(
        Session(), tenant_id=11, job_uuid=str(uuid4())  # type: ignore[arg-type]
    )
    sheet = load_workbook(BytesIO(content))["错误明细"]
    headers = [cell.value for cell in sheet[1]]
    assert headers[:3] == ["原行号", "SKU", "品类"]
    assert sheet.cell(2, 1).value == 9
    assert "ENUM_UNKNOWN" in str(sheet.cell(2, headers.index("错误码") + 1).value)
    assert "Dictionaries" in str(sheet.cell(2, headers.index("修复建议") + 1).value)


def test_ten_thousand_row_preflight_core_performance() -> None:
    started = perf_counter()
    content = _workbook_bytes([
        [f"PERF-{index:05d}", f"性能产品 {index}", "sofa", "active", None, 100, 80, 70, "cm"]
        for index in range(10_000)
    ])
    sheet = load_workbook(BytesIO(content), data_only=False)["Product_Master"]
    headers = [cell.value for cell in sheet[1]]
    mapping = suggest_mapping(headers)
    valid = 0
    for row in sheet.iter_rows(min_row=2, values_only=True):
        source = dict(zip(headers, row))
        values = {code: source[header] for code, header in mapping.items()}
        _, errors, _ = _normalize_values(values, sku_from_cell=values.get("sku"))
        valid += not errors
    elapsed = perf_counter() - started
    assert valid == 10_000
    assert elapsed < 12, f"10k-row preflight core took {elapsed:.2f}s"


async def _seed_inventory(settings, identity: dict[str, object], sku: str) -> None:
    from sqlalchemy import text

    from furniscope_api.database import Database

    database = Database(settings)
    try:
        async with database.session_factory() as session:
            version_id = await session.scalar(text("""
                INSERT INTO forecast_data_versions(
                  tenant_id,created_by,filename,raw_storage_key,raw_sha256,rules,
                  quality,columns_info,canonical_storage_key,canonical_sha256,
                  status,confirmed_at
                ) VALUES(
                  :tenant,:user,'inventory.xlsx','test/inventory.xlsx',:hash,
                  CAST(:rules AS jsonb),'{}'::jsonb,'[]'::jsonb,
                  'test/inventory-canonical.json',:hash,'confirmed',now()
                ) RETURNING id
            """), {
                "tenant": identity["tenant_id"], "user": identity["user_id"],
                "hash": "a" * 64, "rules": json.dumps({"kind": "inventory"}),
            })
            for site, units in (("US", 42), ("DE", 8)):
                await session.execute(text("""
                    INSERT INTO inventory_facts_daily(
                      tenant_id,data_version_id,fact_date,sku,site,
                      inventory_units,source_record_sha256
                    ) VALUES(:tenant,:version,CURRENT_DATE,:sku,:site,:units,:hash)
                """), {
                    "tenant": identity["tenant_id"], "version": version_id,
                    "sku": sku, "site": site, "units": units, "hash": "b" * 64,
                })
            await session.execute(text("""
                INSERT INTO sales_facts_daily(
                  tenant_id,data_version_id,fact_date,sku,site,sales_units,
                  sales_status,source_record_sha256
                ) VALUES(:tenant,:version,CURRENT_DATE,:sku,'US',5,'active',:hash)
            """), {
                "tenant": identity["tenant_id"], "version": version_id,
                "sku": sku, "hash": "c" * 64,
            })
            await session.commit()
    finally:
        await database.close()


async def _projected_fact_skus(settings, identity: dict[str, object]) -> tuple[set[str], set[str]]:
    from sqlalchemy import text

    from furniscope_api.database import Database

    database = Database(settings)
    try:
        async with database.session_factory() as session:
            inventory = await session.execute(text("""
                SELECT DISTINCT sku FROM inventory_facts_daily WHERE tenant_id=:tenant
            """), {"tenant": identity["tenant_id"]})
            sales = await session.execute(text("""
                SELECT DISTINCT sku FROM sales_facts_daily WHERE tenant_id=:tenant
            """), {"tenant": identity["tenant_id"]})
            return (
                {str(row["sku"]) for row in inventory.mappings().all()},
                {str(row["sku"]) for row in sales.mappings().all()},
            )
    finally:
        await database.close()


@pytest.mark.skipif(not TEST_DSN, reason="needs isolated PostgreSQL")
def test_product_catalog_and_transactional_import_http_lifecycle(tmp_path) -> None:
    import asyncpg
    from fastapi.testclient import TestClient

    from furniscope_api.app import create_app
    from test_enterprise_http import fixture_settings, seed_enterprise

    settings = fixture_settings(TEST_DSN, tmp_path)
    identities = asyncio.run(seed_enterprise(settings))
    own_identity, other_identity = identities
    suffix = uuid4().hex[:10].upper()

    with TestClient(create_app(settings)) as client:
        def login(identity):
            response = client.post("/api/v1/auth/login", json={
                "email": identity["email"], "password": identity["password"],
            })
            assert response.status_code == 200, response.text
            return {"Authorization": f"Bearer {response.json()['data']['access_token']}"}

        own = login(own_identity)
        other = login(other_identity)
        create = client.post(
            "/api/v1/products",
            headers={**own, "Idempotency-Key": f"product-{suffix}"},
            json={
                "sku": f"EDIT-{suffix}",
                "name": "可编辑产品",
                "category_code": "table",
                "description": "待清空",
            },
        )
        assert create.status_code == 201, create.text
        product_id = create.json()["data"]["product_id"]
        asyncio.run(_seed_inventory(settings, own_identity, f"EDIT-{suffix}"))
        detail = client.get(f"/api/v1/products/{product_id}", headers=own)
        updated_sku = f"EDITED-{suffix}"
        updated = client.patch(
            f"/api/v1/products/{product_id}",
            headers={**own, "If-Match": detail.headers["etag"]},
            json={
                "sku": updated_sku,
                "category_code": "chair",
                "description": None,
            },
        )
        assert updated.status_code == 200, updated.text
        assert updated.json()["data"]["sku"] == updated_sku
        assert updated.json()["data"]["category_code"] == "chair"
        assert updated.json()["data"]["description"] is None
        inventory_skus, sales_skus = asyncio.run(
            _projected_fact_skus(settings, own_identity)
        )
        assert updated_sku in inventory_skus and f"EDIT-{suffix}" not in inventory_skus
        assert updated_sku in sales_skus and f"EDIT-{suffix}" not in sales_skus

        relations = client.put(
            f"/api/v1/products/{product_id}/relations",
            headers=own,
            json={"items": [
                {
                    "group_type": "spu", "group_code": f"SPU-{suffix}",
                    "group_name": "测试 SPU", "member_role": "variant", "quantity": 1,
                },
                {
                    "group_type": "variant", "group_code": f"VAR-{suffix}",
                    "group_name": "测试变体", "member_role": "parent", "quantity": 1,
                },
                {
                    "group_type": "bundle", "group_code": f"SET-{suffix}",
                    "group_name": "测试套装", "member_role": "item", "quantity": 2,
                },
                {
                    "group_type": "bom", "group_code": f"BOM-{suffix}",
                    "group_name": "测试 BOM", "member_role": "component", "quantity": 0.5,
                },
            ]},
        )
        assert relations.status_code == 200, relations.text
        assert len(relations.json()["data"]["items"]) == 4

        inventory = client.get(
            f"/api/v1/products/{product_id}/inventory-summary", headers=own
        )
        assert inventory.status_code == 200, inventory.text
        inventory_data = inventory.json()["data"]
        assert inventory_data["is_realtime"] is False
        assert inventory_data["total_inventory_units"] == "50.000000"
        assert inventory_data["import_status"]["status"] == "confirmed"

        template = client.get("/api/v1/product-imports/template", headers=own)
        assert template.status_code == 200
        assert template.headers["x-template-version"] == TEMPLATE_VERSION
        assert template.headers["x-content-sha256"] == sha256(template.content).hexdigest()
        assert template.headers["x-schema-sha256"] == template_schema_sha256()

        import_skus = [f"IMP-{suffix}-1", f"IMP-{suffix}-2"]
        content = _workbook_bytes([
            [
                import_skus[0], "导入沙发", "sofa", "active", None,
                220, 95, 82, "cm", None, None, None, None, None,
                f"SPU-IMP-{suffix}",
            ],
            [import_skus[1], "导入桌", "table", "sample", None, 120, 60, 42, "cm"],
        ], alias_rows=[["erp", f"ERP-{suffix}", import_skus[0]]])
        idempotency_key = f"import-{suffix}"
        preflight = client.post(
            "/api/v1/product-imports:preflight",
            headers={**own, "Idempotency-Key": idempotency_key},
            files={"file": ("products.xlsx", content, XLSX_MIME)},
            data={"import_mode": "create_only"},
        )
        assert preflight.status_code == 200, preflight.text
        job = preflight.json()["data"]
        assert job["status"] == "ready"
        assert (job["total_rows"], job["valid_rows"], job["warning_rows"]) == (2, 2, 1)

        replay = client.post(
            "/api/v1/product-imports:preflight",
            headers={**own, "Idempotency-Key": idempotency_key},
            files={"file": ("products.xlsx", content, XLSX_MIME)},
            data={"import_mode": "create_only"},
        )
        assert replay.status_code == 200
        assert replay.json()["data"]["job_uuid"] == job["job_uuid"]
        assert client.get(
            f"/api/v1/product-imports/{job['job_uuid']}", headers=other
        ).status_code == 404

        stale = client.post(
            f"/api/v1/product-imports/{job['job_uuid']}:commit",
            headers=own,
            json={"preview_sha256": "0" * 64},
        )
        assert stale.status_code == 409
        committed = client.post(
            f"/api/v1/product-imports/{job['job_uuid']}:commit",
            headers=own,
            json={"preview_sha256": job["preview_sha256"]},
        )
        assert committed.status_code == 200, committed.text
        committed_data = committed.json()["data"]
        assert committed_data["status"] == "completed"
        assert committed_data["imported_rows"] == 2
        assert committed_data["created_rows"] == 2
        assert committed_data["alias_rows"] == 1
        for sku in import_skus:
            listed = client.get(
                "/api/v1/products", headers=own, params={"keyword": sku}
            )
            assert listed.json()["data"]["total"] == 1
        mappings = client.get("/api/v1/forecast/sku-mappings", headers=own)
        assert any(
            item["source_sku"] == f"ERP-{suffix}"
            and item["product_sku"] == import_skus[0]
            for item in mappings.json()["data"]["items"]
        )

        create_only_existing_content = _workbook_bytes([
            [import_skus[0], "不应覆盖的名称", "chair", "active"],
        ], headers=["SKU", "产品名称", "品类代码", "生命周期"])
        create_only_existing = client.post(
            "/api/v1/product-imports:preflight",
            headers={**own, "Idempotency-Key": f"create-only-existing-{suffix}"},
            files={"file": ("create-only-existing.xlsx", create_only_existing_content, XLSX_MIME)},
            data={"import_mode": "create_only"},
        )
        assert create_only_existing.status_code == 200, create_only_existing.text
        create_only_job = create_only_existing.json()["data"]
        assert create_only_job["status"] == "ready"
        assert create_only_job["rows"][0]["planned_action"] == "skip"
        assert "SKU_EXISTS_SKIPPED" in _issue_codes(
            create_only_job["rows"][0]["validation_warnings"]
        )
        create_only_commit = client.post(
            f"/api/v1/product-imports/{create_only_job['job_uuid']}:commit",
            headers=own,
            json={"preview_sha256": create_only_job["preview_sha256"]},
        )
        assert create_only_commit.status_code == 200, create_only_commit.text
        assert create_only_commit.json()["data"]["imported_rows"] == 0
        assert create_only_commit.json()["data"]["skipped_rows"] == 1
        unchanged_import = client.get(
            "/api/v1/products", headers=own, params={"keyword": import_skus[0]}
        ).json()["data"]["items"][0]
        assert unchanged_import["name"] == "导入沙发"

        upsert_content = _workbook_bytes([
            [import_skus[0], "导入沙发更新版", "chair", "active"],
        ], headers=["SKU", "产品名称", "品类代码", "生命周期"])
        upsert = client.post(
            "/api/v1/product-imports:preflight",
            headers={**own, "Idempotency-Key": f"upsert-{suffix}"},
            files={"file": ("upsert.xlsx", upsert_content, XLSX_MIME)},
            data={"import_mode": "upsert"},
        )
        assert upsert.status_code == 200, upsert.text
        upsert_job = upsert.json()["data"]
        assert upsert_job["status"] == "ready"
        assert upsert_job["rows"][0]["planned_action"] == "update"
        upsert_commit = client.post(
            f"/api/v1/product-imports/{upsert_job['job_uuid']}:commit",
            headers=own,
            json={"preview_sha256": upsert_job["preview_sha256"]},
        )
        assert upsert_commit.status_code == 200, upsert_commit.text
        assert upsert_commit.json()["data"]["updated_rows"] == 1
        updated_import = client.get(
            "/api/v1/products", headers=own, params={"keyword": import_skus[0]}
        ).json()["data"]["items"][0]
        assert updated_import["name"] == "导入沙发更新版"
        assert updated_import["category_code"] == "chair"
        assert updated_import["analysis_status"] == "draft"

        duplicate_sku = f"DUP-{suffix}"
        blocked_content = _workbook_bytes([
            [duplicate_sku, "重复一", "sofa"],
            [duplicate_sku.lower(), "重复二", "sofa"],
        ], headers=["SKU", "产品名称", "品类代码"])
        blocked = client.post(
            "/api/v1/product-imports:preflight",
            headers={**own, "Idempotency-Key": f"blocked-{suffix}"},
            files={"file": ("blocked.xlsx", blocked_content, XLSX_MIME)},
        )
        assert blocked.status_code == 200, blocked.text
        blocked_job = blocked.json()["data"]
        assert blocked_job["status"] == "blocked"
        assert blocked_job["error_rows"] == 2
        report = client.get(
            f"/api/v1/product-imports/{blocked_job['job_uuid']}/error-report",
            headers=own,
        )
        assert report.status_code == 200
        report_sheet = load_workbook(BytesIO(report.content))["错误明细"]
        assert "SKU_DUPLICATE_FILE" in " ".join(
            str(cell.value or "") for cell in report_sheet[2]
        )
        refused = client.post(
            f"/api/v1/product-imports/{blocked_job['job_uuid']}:commit",
            headers=own,
            json={"preview_sha256": blocked_job["preview_sha256"]},
        )
        assert refused.status_code == 422
        assert client.get(
            "/api/v1/products", headers=own, params={"keyword": duplicate_sku}
        ).json()["data"]["total"] == 0
        corrected = client.patch(
            f"/api/v1/product-imports/{blocked_job['job_uuid']}/rows/3",
            headers=own,
            json={"normalized_values": {
                "sku": f"{duplicate_sku}-FIXED",
                "name": "重复二已修正",
                "category_code": "sofa",
                "lifecycle_status": "active",
            }},
        )
        assert corrected.status_code == 200, corrected.text
        assert corrected.json()["data"]["status"] == "ready"
        cancelled = client.post(
            f"/api/v1/product-imports/{blocked_job['job_uuid']}:cancel",
            headers=own,
        )
        assert cancelled.status_code == 200
        assert cancelled.json()["data"]["status"] == "cancelled"
        recent = client.get("/api/v1/product-imports/recent?limit=10", headers=own)
        assert any(
            item["job_uuid"] == blocked_job["job_uuid"]
            and item["status"] == "cancelled"
            for item in recent.json()["data"]["items"]
        )

        service_perf_content = _workbook_bytes([
            [f"HTTP-PERF-{suffix}-{index:05d}", f"服务性能产品 {index}", "sofa"]
            for index in range(10_000)
        ], headers=["SKU", "产品名称", "品类代码"])
        service_perf_started = perf_counter()
        service_perf = client.post(
            "/api/v1/product-imports:preflight",
            headers={**own, "Idempotency-Key": f"perf-{suffix}"},
            files={"file": ("performance-10000.xlsx", service_perf_content, XLSX_MIME)},
        )
        service_perf_elapsed = perf_counter() - service_perf_started
        assert service_perf.status_code == 200, service_perf.text
        service_perf_job = service_perf.json()["data"]
        assert service_perf_job["status"] == "ready"
        assert service_perf_job["total_rows"] == 10_000
        assert service_perf_elapsed < 20, (
            f"10k-row service preflight took {service_perf_elapsed:.2f}s"
        )
        assert client.post(
            f"/api/v1/product-imports/{service_perf_job['job_uuid']}:cancel",
            headers=own,
        ).status_code == 200

        latest = client.get(f"/api/v1/products/{product_id}", headers=own)
        archived = client.delete(
            f"/api/v1/products/{product_id}",
            headers={**own, "If-Match": latest.headers["etag"]},
        )
        assert archived.status_code == 200, archived.text
        assert archived.json()["data"]["analysis_status"] == "archived"
        assert client.get(f"/api/v1/products/{product_id}", headers=own).status_code == 404

    async def audit_actions() -> list[str]:
        connection = await asyncpg.connect(TEST_DSN)
        try:
            rows = await connection.fetch("""
                SELECT action_code FROM furniscope.audit_logs
                 WHERE tenant_id=$1 AND action_code LIKE 'product.%'
                 ORDER BY id
            """, own_identity["tenant_id"])
            return [row["action_code"] for row in rows]
        finally:
            await connection.close()

    actions = asyncio.run(audit_actions())
    assert {
        "product.update", "product.relations.update", "product.import.preflight",
        "product.import.commit", "product.archive",
    } <= set(actions)
