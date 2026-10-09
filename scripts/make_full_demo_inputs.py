"""Explicitly labelled, isolated inputs for recording real import workflows."""
from pathlib import Path
import csv
from datetime import date, timedelta
from openpyxl import Workbook

folder = Path(__file__).resolve().parents[1] / "artifacts/demo-video/full-workflow/inputs"
folder.mkdir(parents=True, exist_ok=True)
book = Workbook()
sheet = book.active
sheet.title = "产品主档"
sheet.append(["SKU", "产品名称", "品类代码", "生命周期", "产品说明", "长度", "宽度", "高度", "尺寸单位", "MOQ", "出厂价", "币种"])
sheet.append(["DEMO-20261009-SOFA", "演示专用双人沙发", "sofa", "active", "功能录像样本，不用于真实业务", 180, 90, 85, "cm", 30, 210, "USD"])
book.save(folder / "演示专用产品导入.xlsx")
book = Workbook()
sheet = book.active
sheet.title = "演示产品说明"
sheet.append(["问题", "答案", "资料性质"])
sheet.append(["演示专用北美休闲椅的尺寸是什么", "演示专用北美休闲椅 DEMO-20261009-CHAIR 的尺寸为 85 × 90 × 105 cm。", "本地功能演示样本"])
sheet.append(["演示专用北美休闲椅的材质是什么", "演示专用北美休闲椅使用实木框架与可拆洗面料。MOQ 为 50 件。", "本地功能演示样本"])
sheet2 = book.create_sheet("使用说明")
sheet2.append(["范围", "说明"])
sheet2.append(["用途", "仅用于演示文件上传、索引、检索和引用，不作为真实企业事实。"])
book.save(folder / "演示产品知识说明.xlsx")
with (folder / "演示标准销量.csv").open("w") as f:
    writer = csv.writer(f)
    writer.writerow(["date", "sku", "site", "sales"])
    for day in range(120):
        writer.writerow([(date(2026, 6, 1) + timedelta(days=day)).isoformat(), "DEMO-20261009-CHAIR", "US", 10 + day % 7])
print(folder)
